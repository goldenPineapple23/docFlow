"""
Reading a catalog or customer-list file into a plain table (CLAUDE.md
Section 7.15.2 Steps 4-5; DECISIONS.md D-108).

**Worker only.** This opens a file a prospect sent -- untrusted input -- so it
runs where every other parser runs (Section 7.11), and
`apps/api/tests/test_parsing_boundary.py` forbids importing it in the API.
The file has already passed `file_types.validate_upload` (magic bytes, size,
decompression limits) before it reaches here.

The output is text and nothing else: a header row and data rows, every cell
a string. A spreadsheet number is rendered through `format_cell_value`-style
Decimal formatting, never `str(float)`, so a SKU like 1002 stays "1002" and a
price never becomes 47.500000000000004 (Section 7.1).
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation

# Defined with the rest of the import rules, which the API may import too.
from docflow_core.catalog_import import TABLE_FORMATS  # noqa: E402

MAX_ROWS = 50_000
MAX_COLUMNS = 100
MAX_CELL_CHARS = 2_000


class ImportParseError(Exception):
    """Carries an error-catalog code (IMP-0xx)."""

    def __init__(self, code: str, detail: str):
        super().__init__(f"{code}: {detail}")
        self.code = code
        self.detail = detail  # for logs only: counts and formats, never cell text


@dataclass(frozen=True)
class ParsedTable:
    columns: list[str]
    # Data rows exactly as they sit under the header, blank rows included, so
    # data row i is spreadsheet row `header_row_number + 1 + i`. The report's
    # row numbers are the ones the founder sees in their spreadsheet.
    rows: list[list[str]]
    header_row_number: int


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        try:
            number = Decimal(repr(value)).normalize()
        except InvalidOperation:
            return str(value)
        if number == number.to_integral_value():
            number = number.quantize(Decimal(1))
        return format(number, "f")
    if isinstance(value, int):
        return str(value)
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def _table_from_rows(raw_rows: list[list[str]]) -> ParsedTable:
    rows = [[cell[:MAX_CELL_CHARS] for cell in row] for row in raw_rows]

    def blank(row: list[str]) -> bool:
        return not any(cell.strip() for cell in row)

    header_index = next((i for i, row in enumerate(rows) if not blank(row)), None)
    if header_index is None:
        raise ImportParseError("IMP-002", "no non-empty rows")
    header = rows[header_index]
    data = rows[header_index + 1 :]
    while data and blank(data[-1]):
        data.pop()
    width = len(header)
    while width and not header[width - 1].strip():
        width -= 1  # trailing unnamed columns
    if width > MAX_COLUMNS:
        raise ImportParseError("IMP-003", f"{width} columns")
    if len(data) > MAX_ROWS:
        raise ImportParseError("IMP-003", f"{len(data)} rows")
    if not data:
        raise ImportParseError("IMP-002", "header row only")
    # Every row the same width as the header: short rows padded, extra cells
    # beyond the last named column dropped.
    normalized = [(row + [""] * width)[:width] for row in data]
    return ParsedTable(
        columns=[c.strip() for c in header[:width]],
        rows=normalized,
        header_row_number=header_index + 1,
    )


def _parse_csv(content: bytes) -> ParsedTable:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text_value = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover -- cp1252 decodes almost anything
        raise ImportParseError("IMP-004", "undecodable text")
    sample = text_value[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text_value, newline=""), dialect)
    return _table_from_rows([list(row) for row in reader])


def _parse_xlsx(content: bytes) -> ParsedTable:
    from openpyxl import load_workbook

    # data_only: a formula's cached value, never the formula. keep_vba off:
    # macros are never loaded (Section 7.11: "No active content, ever").
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_vba=False)
    try:
        sheet = workbook.worksheets[0]
        raw = [[cell_text(v) for v in row] for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    return _table_from_rows(raw)


def _parse_xls(content: bytes) -> ParsedTable:
    import xlrd

    book = xlrd.open_workbook(file_contents=content)
    sheet = book.sheet_by_index(0)
    raw = [[cell_text(sheet.cell_value(r, c)) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
    return _table_from_rows(raw)


def parse_table(content: bytes, file_type: str) -> ParsedTable:
    """The first sheet (or the whole CSV) as text. Raises ImportParseError."""
    if file_type not in TABLE_FORMATS:
        raise ImportParseError("IMP-001", f"file type {file_type}")
    try:
        if file_type in ("csv", "txt"):
            return _parse_csv(content)
        if file_type in ("xlsx", "xlsm"):
            return _parse_xlsx(content)
        return _parse_xls(content)
    except ImportParseError:
        raise
    except Exception as exc:  # noqa: BLE001 -- any library failure on a hostile file
        raise ImportParseError("IMP-004", f"{file_type} unreadable: {type(exc).__name__}") from exc
