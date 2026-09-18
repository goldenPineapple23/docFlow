"""
Export files from the approved snapshot (CLAUDE.md Section 7.3 / 7.4, Phase 4).

The rules this module exists to keep:

  * **From the snapshot, never the live tables.** Every function here takes
    the `document_snapshots.snapshot` a person approved and nothing else. It
    has no database access at all, which is the simplest way to make that
    impossible to get wrong.
  * **Round-trip, every time.** `build_export` renders the file, parses it
    back with this module's own reader, and refuses to return it unless the
    result equals the approved data. Section 7.4 makes this a CI test; it is
    also a runtime check, so a file that fails it never reaches a customer --
    that is `EXP-004`.
  * **Deterministic.** The same snapshot always produces byte-identical
    bytes: fixed column order, fixed line endings, sorted JSON keys, and an
    .xlsx whose zip timestamps and document properties are pinned.
  * **Money stays a string.** No value is ever converted to a float. Numbers
    are written exactly as the snapshot holds them ("47.5000"), and the .xlsx
    stores them as text cells, because an Excel number cell is a binary float.

Formats and their columns are decided in DECISIONS.md D-099. CSV and Excel
are flat: one row per line item with the header fields repeated on every row
(Section 3). JSON is nested. IIF is QuickBooks Desktop's import format, as an
Estimate (the founder's choice -- it posts nothing to the books).

Imported only by the worker and tests: rendering and re-reading .xlsx needs
openpyxl, which the API process never loads (Section 7.11,
`apps/api/tests/test_parsing_boundary.py`).
"""

from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Callable

FORMATS: tuple[str, ...] = ("csv", "xlsx", "json", "iif")

MEDIA_TYPES: dict[str, str] = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "json": "application/json",
    # Served as a download, never rendered; QuickBooks reads it from disk.
    "iif": "application/octet-stream",
}

EXTENSIONS: dict[str, str] = {"csv": "csv", "xlsx": "xlsx", "json": "json", "iif": "iif"}

# ── What an export contains ─────────────────────────────────────────────────
#
# Column names are the extraction schema's own field names (Section 8.1), so
# CSV, Excel and JSON describe the same order the same way. `catalog_sku`
# comes first among the line fields: it is the tenant's own SKU, which is
# what their system imports against (D-099). `sku` is what the buyer's PO
# printed.
HEADER_COLUMNS: tuple[str, ...] = (
    "po_number",
    "order_date",
    "requested_delivery_date",
    "buyer_name",
    "buyer_contact_email",
    "ship_to_address",
    "payment_terms",
    "currency",
    "order_total",
    "notes",
)
LINE_COLUMNS: tuple[str, ...] = (
    "line_number",
    "catalog_sku",
    "sku",
    "description",
    "quantity",
    "unit",
    "unit_price",
    "line_total",
)
FLAT_COLUMNS: tuple[str, ...] = HEADER_COLUMNS + LINE_COLUMNS

# QuickBooks Desktop's IIF has no column for these. They are left out of the
# .iif file, and the round-trip comparison for IIF leaves them out too --
# explicitly, here, rather than by the comparison quietly passing (D-099).
# `notes` is here by choice: IIF cannot hold a line break, notes routinely
# have several lines, and failing every such export would make the format
# unusable. The notes stay in the CSV, Excel and JSON exports.
IIF_OMITTED_HEADER: frozenset[str] = frozenset({"buyer_contact_email", "currency", "notes"})
IIF_OMITTED_LINE: frozenset[str] = frozenset({"line_number", "sku", "unit"})

# The account an IIF Estimate is recorded against. "Estimates" is the
# non-posting account QuickBooks Desktop keeps for them; the line account is
# the income account a line would post to if the estimate became an invoice.
# Constants, in one place, pending UAT TC-26 against a real QuickBooks file.
IIF_ESTIMATE_ACCOUNT = "Estimates"
IIF_LINE_ACCOUNT = "Sales"
IIF_ADDRESS_LINES = 5

# Pinned so the same snapshot gives the same .xlsx bytes on any day.
_FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
_FIXED_DOC_TIME = datetime(2000, 1, 1, 0, 0, 0)


class ExportError(Exception):
    """A file that cannot be produced. Always carries an error-catalog code."""

    def __init__(self, code: str, detail: str):
        super().__init__(f"{code}: {detail}")
        self.code = code
        # For logs only: a field NAME or a count, never a document value
        # (Section 7.10).
        self.detail = detail


@dataclass(frozen=True)
class BuiltExport:
    content: bytes
    media_type: str
    extension: str


# ── The order, as every format sees it ──────────────────────────────────────


def order_view(snapshot: dict[str, Any]) -> dict[str, Any]:
    """
    The exportable part of an approved snapshot: the header fields and line
    fields above, in a fixed order, and nothing else.

    `catalog_sku` is read with `.get` because snapshots approved before D-099
    do not carry it; those export with an empty catalog SKU until re-approved.
    """
    header = snapshot.get("header") or {}
    return {
        "header": {name: header.get(name) for name in HEADER_COLUMNS},
        "lines": [
            {
                name: (int(line["line_number"]) if name == "line_number" else line.get(name))
                for name in LINE_COLUMNS
            }
            for line in snapshot.get("lines") or []
        ],
    }


def comparable_view(view: dict[str, Any], fmt: str) -> dict[str, Any]:
    """What a parsed-back file must equal. Everything, except IIF's omissions."""
    if fmt != "iif":
        return view
    return {
        "header": {k: v for k, v in view["header"].items() if k not in IIF_OMITTED_HEADER},
        "lines": [
            {k: v for k, v in line.items() if k not in IIF_OMITTED_LINE} for line in view["lines"]
        ],
    }


# ── Flat rows (CSV and Excel share these) ───────────────────────────────────


def _flat_rows(view: dict[str, Any]) -> list[list[str | None]]:
    header = [view["header"][name] for name in HEADER_COLUMNS]
    if not view["lines"]:
        # An order with no lines still exports its header, once.
        return [header + [None] * len(LINE_COLUMNS)]
    rows = []
    for line in view["lines"]:
        values = [None if line[name] is None else str(line[name]) for name in LINE_COLUMNS]
        rows.append(header + values)
    return rows


def _view_from_flat_rows(rows: list[list[str | None]]) -> dict[str, Any]:
    if not rows:
        raise ExportError("EXP-004", "no data rows")
    width = len(HEADER_COLUMNS)
    header_values = rows[0][:width]
    for row in rows:
        if len(row) != len(FLAT_COLUMNS):
            raise ExportError("EXP-004", "row width")
        if row[:width] != header_values:
            raise ExportError("EXP-004", "header fields differ between rows")

    lines = []
    for row in rows:
        values = dict(zip(LINE_COLUMNS, row[width:]))
        if values["line_number"] is None:
            if len(rows) != 1 or any(v is not None for v in values.values()):
                raise ExportError("EXP-004", "row without a line number")
            continue
        values["line_number"] = int(values["line_number"])
        lines.append(values)
    return {"header": dict(zip(HEADER_COLUMNS, header_values)), "lines": lines}


# ── CSV ─────────────────────────────────────────────────────────────────────
#
# A value that starts with = + - @ (or a tab or carriage return) is run as a
# formula when the file is opened in Excel -- and every value here came from a
# stranger's document (Section 7.2). Such values are written with a leading
# apostrophe, the OWASP-recommended neutraliser. A value that already starts
# with an apostrophe gets one too, so reading the file back is exact: strip
# one leading apostrophe, always. Plain numbers ("-5.00") are left alone.

_FORMULA_TRIGGERS = ("=", "+", "-", "@", "\t", "\r", "'")
_PLAIN_NUMBER = re.compile(r"-?\d+(?:\.\d+)?\Z")


def _csv_escape(value: str | None) -> str:
    if value is None:
        return ""
    if value.startswith(_FORMULA_TRIGGERS) and not _PLAIN_NUMBER.match(value):
        return "'" + value
    return value


def _csv_unescape(value: str) -> str | None:
    if value == "":
        return None
    if value.startswith("'"):
        return value[1:]
    return value


def render_csv(view: dict[str, Any]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(FLAT_COLUMNS)
    for row in _flat_rows(view):
        writer.writerow([_csv_escape(value) for value in row])
    # UTF-8 with a byte-order mark: without it Excel opens "Café" as "CafÃ©".
    return buffer.getvalue().encode("utf-8-sig")


def parse_csv(content: bytes) -> dict[str, Any]:
    reader = csv.reader(io.StringIO(content.decode("utf-8-sig"), newline=""))
    rows = list(reader)
    if not rows or tuple(rows[0]) != FLAT_COLUMNS:
        raise ExportError("EXP-004", "csv column row")
    return _view_from_flat_rows([[_csv_unescape(v) for v in row] for row in rows[1:]])


# ── Excel ───────────────────────────────────────────────────────────────────


def render_xlsx(view: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Order"
    sheet.append(list(FLAT_COLUMNS))
    for row_index, row in enumerate(_flat_rows(view), start=2):
        for column_index, value in enumerate(row, start=1):
            if value is None:
                continue
            cell = sheet.cell(row=row_index, column=column_index)
            cell.value = value
            # Text, always: never a formula (openpyxl would treat "=..." as
            # one) and never a float (an Excel number cell is binary).
            cell.data_type = "s"
            cell.number_format = "@"
    for column_index, name in enumerate(FLAT_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(column_index)].width = max(12, len(name) + 2)
    sheet.freeze_panes = "A2"

    properties = workbook.properties
    properties.creator = "DocFlow"
    properties.lastModifiedBy = None
    properties.created = _FIXED_DOC_TIME
    properties.modified = _FIXED_DOC_TIME

    raw = io.BytesIO()
    workbook.save(raw)
    return _pin_zip(raw.getvalue())


_CORE_XML_TIMESTAMP = re.compile(
    rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:(?:created|modified)>)"
)


def _pin_zip(content: bytes) -> bytes:
    """
    Rewrite an .xlsx so nothing in it depends on when it was made: every zip
    entry gets the same timestamp and is stored uncompressed, and the
    document's created/modified
    properties are pinned whatever the library wrote.
    """
    source = zipfile.ZipFile(io.BytesIO(content))
    output = io.BytesIO()
    # Stored, not deflated: compressed bytes depend on the zlib build, so the
    # same workbook deflated on Windows and on Linux need not match. An order
    # file is a few kilobytes either way.
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_STORED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = _CORE_XML_TIMESTAMP.sub(rb"\g<1>2000-01-01T00:00:00Z\g<2>", data)
            pinned = zipfile.ZipInfo(info.filename, date_time=_FIXED_ZIP_TIME)
            pinned.compress_type = zipfile.ZIP_STORED
            pinned.external_attr = 0o600 << 16
            target.writestr(pinned, data)
    return output.getvalue()


def parse_xlsx(content: bytes) -> dict[str, Any]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    try:
        sheet = workbook["Order"]
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    if not rows or tuple(rows[0]) != FLAT_COLUMNS:
        raise ExportError("EXP-004", "xlsx column row")
    data_rows: list[list[str | None]] = []
    for row in rows[1:]:
        padded = list(row) + [None] * (len(FLAT_COLUMNS) - len(row))
        for value in padded:
            if value is not None and not isinstance(value, str):
                raise ExportError("EXP-004", "xlsx cell is not text")
        data_rows.append(padded)
    return _view_from_flat_rows(data_rows)


# ── JSON ────────────────────────────────────────────────────────────────────

JSON_FORMAT_NAME = "docflow.order"
JSON_FORMAT_VERSION = 1


def render_json(view: dict[str, Any], *, document_id: str, snapshot_hash: str) -> bytes:
    payload = {
        "format": JSON_FORMAT_NAME,
        "version": JSON_FORMAT_VERSION,
        "document_id": document_id,
        "snapshot_hash": snapshot_hash,
        "header": view["header"],
        "lines": view["lines"],
    }
    return (json.dumps(payload, indent=2, sort_keys=True, ensure_ascii=False) + "\n").encode("utf-8")


def parse_json(content: bytes) -> dict[str, Any]:
    payload = json.loads(content.decode("utf-8"))
    if payload.get("format") != JSON_FORMAT_NAME:
        raise ExportError("EXP-004", "json format marker")
    return {"header": payload["header"], "lines": payload["lines"]}


# ── IIF (QuickBooks Desktop) ────────────────────────────────────────────────

IIF_TRNS_COLUMNS: tuple[str, ...] = (
    "TRNSID",
    "TRNSTYPE",
    "DATE",
    "ACCNT",
    "NAME",
    "AMOUNT",
    "DOCNUM",
    "PONUM",
    "TERMS",
    "SHIPDATE",
    "MEMO",
) + tuple(f"SADDR{n}" for n in range(1, IIF_ADDRESS_LINES + 1))
IIF_SPL_COLUMNS: tuple[str, ...] = (
    "SPLID",
    "TRNSTYPE",
    "DATE",
    "ACCNT",
    "NAME",
    "AMOUNT",
    "QNTY",
    "PRICE",
    "INVITEM",
    "MEMO",
)
_IIF_TYPE = "ESTIMATE"
_ISO_DATE = re.compile(r"(\d{4})-(\d{2})-(\d{2})\Z")
_IIF_DATE = re.compile(r"(\d{2})/(\d{2})/(\d{4})\Z")


def _iif_field(value: str | None, name: str) -> str:
    if value is None:
        return ""
    if "\t" in value or "\r" in value or "\n" in value:
        raise ExportError("EXP-005", f"{name} contains a tab or line break")
    return value


def _iif_date(value: str | None, name: str) -> str:
    if value is None:
        return ""
    match = _ISO_DATE.match(value)
    if not match:
        raise ExportError("EXP-005", f"{name} is not a YYYY-MM-DD date")
    year, month, day = match.groups()
    return f"{month}/{day}/{year}"


def _iso_from_iif(value: str) -> str | None:
    if value == "":
        return None
    match = _IIF_DATE.match(value)
    if not match:
        raise ExportError("EXP-004", "iif date")
    month, day, year = match.groups()
    return f"{year}-{month}-{day}"


def _negate(value: str | None) -> str:
    """Flip a sign by text, so "47.50" stays exactly "47.50" (no float, no
    Decimal re-rendering). QuickBooks records estimate lines as negatives."""
    if value is None:
        return ""
    return value[1:] if value.startswith("-") else "-" + value


def _un_negate(value: str) -> str | None:
    if value == "":
        return None
    return value[1:] if value.startswith("-") else "-" + value


def _iif_address(value: str | None) -> list[str]:
    if value is None:
        return [""] * IIF_ADDRESS_LINES
    parts = value.split("\n")
    if len(parts) > IIF_ADDRESS_LINES:
        raise ExportError("EXP-005", "ship_to_address has more than five lines")
    if parts[-1] == "" or any("\r" in part or "\t" in part for part in parts):
        raise ExportError("EXP-005", "ship_to_address has a blank last line or a tab")
    return parts + [""] * (IIF_ADDRESS_LINES - len(parts))


def _check_iif_balances(view: dict[str, Any]) -> None:
    """
    QuickBooks only imports a transaction whose lines add up to its total,
    and needs a customer name and a date. An order missing any would be rejected
    on import -- so no file is produced, and the reviewer is told why
    (EXP-006) instead of the customer finding out inside QuickBooks.
    """
    header = view["header"]
    if not header["buyer_name"]:
        raise ExportError("EXP-006", "buyer_name is empty")
    if not header["order_date"]:
        raise ExportError("EXP-006", "order_date is empty")
    try:
        total = Decimal(header["order_total"]) if header["order_total"] is not None else None
        line_totals = [
            Decimal(line["line_total"]) for line in view["lines"] if line["line_total"] is not None
        ]
    except InvalidOperation as exc:
        raise ExportError("EXP-006", "a total is not a number") from exc
    if total is None or len(line_totals) != len(view["lines"]) or not view["lines"]:
        raise ExportError("EXP-006", "order total or a line total is missing")
    if sum(line_totals, Decimal("0")) != total:
        raise ExportError("EXP-006", "line totals do not add up to the order total")


def render_iif(view: dict[str, Any]) -> bytes:
    _check_iif_balances(view)
    header = view["header"]
    date = _iif_date(header["order_date"], "order_date")
    name = _iif_field(header["buyer_name"], "buyer_name")

    trns = [
        "",
        _IIF_TYPE,
        date,
        IIF_ESTIMATE_ACCOUNT,
        name,
        _iif_field(header["order_total"], "order_total"),
        "",  # DOCNUM: left to QuickBooks, which numbers estimates itself.
        _iif_field(header["po_number"], "po_number"),
        _iif_field(header["payment_terms"], "payment_terms"),
        _iif_date(header["requested_delivery_date"], "requested_delivery_date"),
        "",  # MEMO: notes are not carried -- see IIF_OMITTED_HEADER.
        *_iif_address(header["ship_to_address"]),
    ]
    lines = [
        "!TRNS\t" + "\t".join(IIF_TRNS_COLUMNS),
        "!SPL\t" + "\t".join(IIF_SPL_COLUMNS),
        "!ENDTRNS",
        "TRNS\t" + "\t".join(trns),
    ]
    for line in view["lines"]:
        spl = [
            "",
            _IIF_TYPE,
            date,
            IIF_LINE_ACCOUNT,
            "",
            _negate(_iif_field(line["line_total"], "line_total")),
            _negate(_iif_field(line["quantity"], "quantity")),
            _iif_field(line["unit_price"], "unit_price"),
            _iif_field(line["catalog_sku"], "catalog_sku"),
            _iif_field(line["description"], "description"),
        ]
        lines.append("SPL\t" + "\t".join(spl))
    lines.append("ENDTRNS")
    text_value = "\r\n".join(lines) + "\r\n"
    try:
        # QuickBooks Desktop reads IIF as Windows-1252, not UTF-8.
        return text_value.encode("cp1252")
    except UnicodeEncodeError as exc:
        raise ExportError("EXP-005", "a character QuickBooks' file format can't hold") from exc


def parse_iif(content: bytes) -> dict[str, Any]:
    rows = content.decode("cp1252").split("\r\n")
    if rows[-1] != "":
        raise ExportError("EXP-004", "iif missing final line break")
    rows = rows[:-1]
    expected_heads = [
        "!TRNS\t" + "\t".join(IIF_TRNS_COLUMNS),
        "!SPL\t" + "\t".join(IIF_SPL_COLUMNS),
        "!ENDTRNS",
    ]
    if rows[:3] != expected_heads or rows[-1] != "ENDTRNS" or not rows[3].startswith("TRNS\t"):
        raise ExportError("EXP-004", "iif structure")

    trns = dict(zip(IIF_TRNS_COLUMNS, rows[3].split("\t")[1:]))
    address = [trns[f"SADDR{n}"] for n in range(1, IIF_ADDRESS_LINES + 1)]
    while address and address[-1] == "":
        address.pop()

    def blank_to_none(value: str) -> str | None:
        return value if value != "" else None

    header = {
        "po_number": blank_to_none(trns["PONUM"]),
        "order_date": _iso_from_iif(trns["DATE"]),
        "requested_delivery_date": _iso_from_iif(trns["SHIPDATE"]),
        "buyer_name": blank_to_none(trns["NAME"]),
        "ship_to_address": "\n".join(address) if address else None,
        "payment_terms": blank_to_none(trns["TERMS"]),
        "order_total": blank_to_none(trns["AMOUNT"]),
    }
    lines = []
    for row in rows[4:-1]:
        if not row.startswith("SPL\t"):
            raise ExportError("EXP-004", "iif line row")
        spl = dict(zip(IIF_SPL_COLUMNS, row.split("\t")[1:]))
        lines.append(
            {
                "catalog_sku": blank_to_none(spl["INVITEM"]),
                "description": blank_to_none(spl["MEMO"]),
                "quantity": _un_negate(spl["QNTY"]),
                "unit_price": blank_to_none(spl["PRICE"]),
                "line_total": _un_negate(spl["AMOUNT"]),
            }
        )
    return {
        "header": {name: header[name] for name in HEADER_COLUMNS if name not in IIF_OMITTED_HEADER},
        "lines": [
            {name: line[name] for name in LINE_COLUMNS if name not in IIF_OMITTED_LINE}
            for line in lines
        ],
    }


# ── The one entry point ─────────────────────────────────────────────────────

PARSERS: dict[str, Callable[[bytes], dict[str, Any]]] = {
    "csv": parse_csv,
    "xlsx": parse_xlsx,
    "json": parse_json,
    "iif": parse_iif,
}


def render(snapshot: dict[str, Any], snapshot_hash: str, fmt: str) -> bytes:
    view = order_view(snapshot)
    if fmt == "csv":
        return render_csv(view)
    if fmt == "xlsx":
        return render_xlsx(view)
    if fmt == "json":
        return render_json(view, document_id=str(snapshot["document_id"]), snapshot_hash=snapshot_hash)
    if fmt == "iif":
        return render_iif(view)
    raise ExportError("EXP-002", f"unknown format {fmt!r}")


def build_export(snapshot: dict[str, Any], snapshot_hash: str, fmt: str) -> BuiltExport:
    """
    The file for one approved snapshot, verified before it is returned.

    Two checks, both Section 7.4, both at runtime as well as in CI:
      * it is rendered twice and the bytes must match (determinism);
      * it is parsed back and must equal the approved data (round trip).
    Either failing is `EXP-004`: the file is not given to anyone.
    """
    content = render(snapshot, snapshot_hash, fmt)
    if render(snapshot, snapshot_hash, fmt) != content:
        raise ExportError("EXP-004", "two renderings differ")
    try:
        parsed = PARSERS[fmt](content)
    except ExportError:
        raise
    except Exception as exc:  # noqa: BLE001 -- any failure to read it back is a failed check
        raise ExportError("EXP-004", f"file could not be read back: {type(exc).__name__}") from exc
    if parsed != comparable_view(order_view(snapshot), fmt):
        raise ExportError("EXP-004", "parsed file differs from the approved snapshot")
    return BuiltExport(content=content, media_type=MEDIA_TYPES[fmt], extension=EXTENSIONS[fmt])
