"""
Export files (CLAUDE.md Section 7.4, Phase 4).

The Phase 4 exit criteria, as tests:
  * "the exported file, parsed back, equals the approved data exactly" --
    for all four formats, including the values most likely to break a file
    (commas, quotes, line breaks, non-English text, formula-looking text);
  * "exporting twice produces byte-identical files" -- asserted two ways:
    rendering twice in one process, and against pinned SHA-256 digests, so a
    library upgrade that changes the bytes is a visible, deliberate diff.

All data is fictional (CLAUDE.md Section 0 rule 4).
"""

from __future__ import annotations

import copy
import hashlib
import io
import zipfile

import pytest

from docflow_core import exports
from docflow_core.errors import CATALOG
from docflow_core.exports import ExportError, build_export, comparable_view, order_view

SNAPSHOT_HASH = "0" * 64


def _snapshot() -> dict:
    return {
        "document_id": "11111111-1111-1111-1111-111111111111",
        "header": {
            "po_number": "ACME-2291",
            "order_date": "2025-05-28",
            "requested_delivery_date": "2025-06-04",
            "buyer_name": "Acme Test Café, \"West\" Branch",
            "buyer_contact_email": "orders@acme-test.example",
            "ship_to_address": "1 Test Street\nSuite 2\nTestville, TS 00000",
            "payment_terms": "Net 30",
            "currency": "USD",
            "order_total": "942.00",
            "notes": "Deliver to the back door.\nCall ahead.",
        },
        "lines": [
            {
                "line_number": 1,
                "catalog_sku": "TEST-1001",
                "sku": "B-77",
                "description": "Test Beans, 5lb",
                "quantity": "12.0000",
                "unit": "CS",
                "unit_price": "47.5000",
                "line_total": "570.00",
                "matched_item_id": "22222222-2222-2222-2222-222222222222",
                "matched_uom": "CS",
            },
            {
                "line_number": 2,
                "catalog_sku": None,
                "sku": None,
                "description": "Test Syrup 750ml",
                "quantity": "6.0000",
                "unit": "CS",
                "unit_price": "62.0000",
                "line_total": "372.00",
                "matched_item_id": None,
                "matched_uom": None,
            },
        ],
    }


# ── Round trip: the Phase 4 exit criterion ──────────────────────────────────


@pytest.mark.parametrize("fmt", exports.FORMATS)
def test_every_format_parses_back_to_exactly_the_approved_data(fmt):
    snapshot = _snapshot()
    built = build_export(snapshot, SNAPSHOT_HASH, fmt)

    parsed = exports.PARSERS[fmt](built.content)

    assert parsed == comparable_view(order_view(snapshot), fmt)


@pytest.mark.parametrize("fmt", ["csv", "xlsx", "json"])
def test_the_flat_and_json_formats_carry_every_field(fmt):
    """Only IIF may leave fields out, and only the ones it names."""
    snapshot = _snapshot()
    parsed = exports.PARSERS[fmt](build_export(snapshot, SNAPSHOT_HASH, fmt).content)
    assert parsed == order_view(snapshot)
    assert set(parsed["header"]) == set(exports.HEADER_COLUMNS)


def test_iif_leaves_out_exactly_the_fields_it_says_it_does():
    snapshot = _snapshot()
    parsed = exports.parse_iif(build_export(snapshot, SNAPSHOT_HASH, "iif").content)
    assert set(exports.HEADER_COLUMNS) - set(parsed["header"]) == exports.IIF_OMITTED_HEADER
    assert set(exports.LINE_COLUMNS) - set(parsed["lines"][0]) == exports.IIF_OMITTED_LINE


@pytest.mark.parametrize("fmt", ["csv", "xlsx", "json"])
def test_hostile_values_survive_the_round_trip(fmt):
    snapshot = _snapshot()
    snapshot["header"]["po_number"] = '=HYPERLINK("http://example.test","x")'
    snapshot["header"]["notes"] = "'already quoted, with a comma\r\nand a CRLF"
    snapshot["lines"][0]["description"] = "@SUM(A1:A2)"
    snapshot["lines"][0]["sku"] = "-B77"
    snapshot["lines"][1]["description"] = "+Plus, \"quotes\" and ünïcödé"

    parsed = exports.PARSERS[fmt](build_export(snapshot, SNAPSHOT_HASH, fmt).content)

    assert parsed == order_view(snapshot)


def test_an_order_with_no_lines_exports_its_header():
    snapshot = _snapshot()
    snapshot["lines"] = []
    for fmt in ("csv", "xlsx", "json"):
        parsed = exports.PARSERS[fmt](build_export(snapshot, SNAPSHOT_HASH, fmt).content)
        assert parsed == order_view(snapshot)


def test_a_snapshot_approved_before_catalog_skus_exports_with_the_column_empty():
    snapshot = _snapshot()
    for line in snapshot["lines"]:
        del line["catalog_sku"]
    parsed = exports.parse_csv(build_export(snapshot, SNAPSHOT_HASH, "csv").content)
    assert [line["catalog_sku"] for line in parsed["lines"]] == [None, None]


# ── Determinism: "exporting twice produces byte-identical files" ────────────


@pytest.mark.parametrize("fmt", exports.FORMATS)
def test_exporting_twice_is_byte_identical(fmt):
    first = build_export(_snapshot(), SNAPSHOT_HASH, fmt).content
    second = build_export(copy.deepcopy(_snapshot()), SNAPSHOT_HASH, fmt).content
    assert first == second


# Pinned digests for `_snapshot()`. Same snapshot -> same bytes, on any day,
# in any process. A change here means the file a customer downloads has
# changed: update it deliberately, with the reason in the commit.
PINNED_SHA256 = {
    "csv": "7a3950feb812215287b7e8cb2873ddaf9470a233f13d19a6e34deb8d038742bb",
    "xlsx": "ac1ba7e0b9b2579b6f755251dce0eaba07bca3c9b872ba223273e1108555c023",
    "json": "c867befa8fb6b0ecbdd5ae1ef640917b6010f918014af8d7a98f0e6321a35a66",
    "iif": "0f4cad938b2c7ac7d3d9e30528cb0dd4a1da5cbd4c87be7726aa8e7d96cb6cc4",
}


@pytest.mark.parametrize("fmt", exports.FORMATS)
def test_the_bytes_match_the_pinned_digest(fmt):
    content = build_export(_snapshot(), SNAPSHOT_HASH, fmt).content
    assert hashlib.sha256(content).hexdigest() == PINNED_SHA256[fmt]


@pytest.mark.parametrize("platform", ["win32", "linux", "darwin"])
def test_the_bytes_do_not_depend_on_the_operating_system(platform, monkeypatch):
    """CI (Linux) once disagreed with Windows about the .xlsx digest, because
    zipfile records the OS that wrote each entry. Same snapshot, same bytes,
    wherever the worker runs."""
    import sys

    monkeypatch.setattr(sys, "platform", platform)
    for fmt in exports.FORMATS:
        content = build_export(_snapshot(), SNAPSHOT_HASH, fmt).content
        assert hashlib.sha256(content).hexdigest() == PINNED_SHA256[fmt], fmt


def test_the_xlsx_carries_no_timestamp_from_when_it_was_made():
    content = build_export(_snapshot(), SNAPSHOT_HASH, "xlsx").content
    archive = zipfile.ZipFile(io.BytesIO(content))
    assert {info.date_time for info in archive.infolist()} == {(1980, 1, 1, 0, 0, 0)}
    core = archive.read("docProps/core.xml")
    assert b"2000-01-01T00:00:00Z" in core


# ── Structure the formats promise ───────────────────────────────────────────


def test_csv_repeats_the_header_on_every_line_row():
    content = build_export(_snapshot(), SNAPSHOT_HASH, "csv").content.decode("utf-8-sig")
    import csv

    rows = list(csv.reader(io.StringIO(content, newline="")))
    assert tuple(rows[0]) == exports.FLAT_COLUMNS
    assert len(rows) == 3
    assert rows[1][0] == rows[2][0] == "ACME-2291"
    assert rows[1][exports.FLAT_COLUMNS.index("catalog_sku")] == "TEST-1001"


def test_csv_neutralises_formula_text_but_not_negative_numbers():
    snapshot = _snapshot()
    snapshot["header"]["po_number"] = "=cmd"
    snapshot["lines"][0]["line_total"] = "-5.00"
    snapshot["lines"][0]["quantity"] = "-1.0000"
    snapshot["lines"][0]["unit_price"] = "5.0000"
    snapshot["header"]["order_total"] = "367.00"
    text = build_export(snapshot, SNAPSHOT_HASH, "csv").content.decode("utf-8-sig")
    assert "'=cmd" in text
    assert ",-5.00" in text and "'-5.00" not in text


def test_xlsx_stores_formula_text_and_money_as_text_never_formulas_or_floats():
    snapshot = _snapshot()
    snapshot["header"]["po_number"] = "=cmd"
    content = build_export(snapshot, SNAPSHOT_HASH, "xlsx").content
    sheet_xml = zipfile.ZipFile(io.BytesIO(content)).read("xl/worksheets/sheet1.xml")
    assert b"<f>" not in sheet_xml
    # Every populated cell is a string cell: no numeric (<c ... t="n">) or
    # untyped cells, which Excel would read as binary floats.
    assert b't="n"' not in sheet_xml
    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(content))["Order"]
    assert sheet["A2"].value == "=cmd" and sheet["A2"].data_type == "s"
    assert sheet.cell(row=2, column=exports.FLAT_COLUMNS.index("unit_price") + 1).value == "47.5000"


def test_json_is_nested_and_names_the_snapshot_it_came_from():
    import json

    payload = json.loads(build_export(_snapshot(), SNAPSHOT_HASH, "json").content)
    assert payload["format"] == "docflow.order"
    assert payload["snapshot_hash"] == SNAPSHOT_HASH
    assert payload["lines"][0]["unit_price"] == "47.5000"  # a string, never a number


def test_iif_is_a_quickbooks_estimate_with_signed_lines():
    text = build_export(_snapshot(), SNAPSHOT_HASH, "iif").content.decode("cp1252")
    rows = text.split("\r\n")
    assert rows[0].startswith("!TRNS\t")
    trns = rows[3].split("\t")
    assert trns[0] == "TRNS" and trns[2] == "ESTIMATE" and trns[3] == "05/28/2025"
    spl = rows[4].split("\t")
    assert spl[0] == "SPL" and spl[6] == "-570.00" and spl[7] == "-12.0000" and spl[9] == "TEST-1001"
    assert rows[-2] == "ENDTRNS" and rows[-1] == ""


# ── IIF refuses what QuickBooks would reject or garble ──────────────────────


def _iif_error(snapshot: dict) -> str:
    with pytest.raises(ExportError) as excinfo:
        build_export(snapshot, SNAPSHOT_HASH, "iif")
    return excinfo.value.code


def test_iif_refuses_an_order_whose_lines_do_not_add_up():
    snapshot = _snapshot()
    snapshot["header"]["order_total"] = "950.00"  # e.g. freight included
    assert _iif_error(snapshot) == "EXP-006"


def test_iif_refuses_an_order_with_no_order_date():
    """Found driving the real app: an order with no date made an IIF with a
    blank DATE, which QuickBooks would reject on import."""
    snapshot = _snapshot()
    snapshot["header"]["order_date"] = None
    assert _iif_error(snapshot) == "EXP-006"


def test_iif_refuses_an_order_with_no_buyer_name():
    snapshot = _snapshot()
    snapshot["header"]["buyer_name"] = None
    assert _iif_error(snapshot) == "EXP-006"


def test_iif_refuses_a_line_break_it_cannot_hold():
    snapshot = _snapshot()
    snapshot["lines"][0]["description"] = "Two\nlines"
    assert _iif_error(snapshot) == "EXP-005"


def test_iif_refuses_a_character_quickbooks_cannot_read():
    snapshot = _snapshot()
    snapshot["lines"][0]["description"] = "Test beans ☕"  # not in Windows-1252
    assert _iif_error(snapshot) == "EXP-005"


def test_iif_refuses_an_address_longer_than_five_lines():
    snapshot = _snapshot()
    snapshot["header"]["ship_to_address"] = "\n".join(f"Line {n}" for n in range(6))
    assert _iif_error(snapshot) == "EXP-005"


def test_iif_carries_multi_line_notes_by_leaving_them_out():
    snapshot = _snapshot()
    assert "\n" in snapshot["header"]["notes"]
    build_export(snapshot, SNAPSHOT_HASH, "iif")  # does not raise


# ── The runtime integrity check ─────────────────────────────────────────────


def test_a_file_that_does_not_parse_back_to_the_snapshot_is_never_returned(monkeypatch):
    real_parse = exports.PARSERS["csv"]

    def corrupting_parse(content):
        parsed = real_parse(content)
        parsed["lines"][0]["line_total"] = "57.00"  # a dropped digit
        return parsed

    monkeypatch.setitem(exports.PARSERS, "csv", corrupting_parse)
    with pytest.raises(ExportError) as excinfo:
        build_export(_snapshot(), SNAPSHOT_HASH, "csv")
    assert excinfo.value.code == "EXP-004"


def test_a_renderer_that_is_not_deterministic_is_caught(monkeypatch):
    counter = iter(range(100))
    monkeypatch.setattr(exports, "render_json", lambda view, **_: f'{{"n": {next(counter)}}}'.encode())
    with pytest.raises(ExportError) as excinfo:
        build_export(_snapshot(), SNAPSHOT_HASH, "json")
    assert excinfo.value.code == "EXP-004"


def test_an_unknown_format_is_a_catalog_error():
    with pytest.raises(ExportError) as excinfo:
        build_export(_snapshot(), SNAPSHOT_HASH, "pdf")
    assert excinfo.value.code == "EXP-002"


def test_every_code_this_module_raises_is_in_the_catalog():
    import inspect
    import re

    source = inspect.getsource(exports)
    for code in set(re.findall(r'"(EXP-\d{3})"', source)):
        assert code in CATALOG, code
