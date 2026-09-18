"""The preview module's own text extraction (DECISIONS.md D-092)."""

from __future__ import annotations

from io import BytesIO

import docx

from docflow_core import previews
from docflow_core.file_types import FileTypeName


def test_a_word_preview_includes_table_rows():
    document = docx.Document()
    document.add_paragraph("PURCHASE ORDER")
    table = document.add_table(rows=0, cols=3)
    for row in (["SKU", "Description", "Qty"], ["TEST-SKU-1", "Test Widget", "12"]):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
    buffer = BytesIO()
    document.save(buffer)

    preview = previews.build_preview(buffer.getvalue(), FileTypeName.DOCX)

    assert preview is not None
    assert b"TEST-SKU-1\tTest Widget\t12" in preview.content


def test_text_preview_truncates_and_refuses_blank_text():
    assert previews.text_preview("   \n ") is None
    long = previews.text_preview("x" * (previews.MAX_PREVIEW_CHARS + 10))
    assert long is not None and long.content.endswith("[…truncated for viewing]".encode())
