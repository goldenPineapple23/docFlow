"""
The isolated parsing/extraction task (CLAUDE.md Section 7.11: "parsing never
runs in the web process... a worker that dies takes one document to `failed`,
never the app"). This module IS that isolated worker -- the API layer
(apps/api/app/routers/documents.py) only validates bytes and enqueues this
task, it never parses. Tier 2 conversion happens here too, one step earlier
in the same process (see app/conversion.py), because "Tier 2 is where the
risk is -- treat conversion as parsing."

`tenant_id` is a required argument, not derived from anything inside this
task: the enqueuing caller (the upload endpoint) already resolved it from
the authenticated session per CLAUDE.md Section 7.5, and passes it through
so this task can open a correctly-scoped `tenant_session()` from the start --
there is no query in this module that runs before a tenant context is set.

Every DB write goes through `docflow_core.db.tenant_session()` -- never a
raw connection, never `admin_data_access` (Section 7.15.1's import boundary
doesn't apply here; this module is not `/admin/*` and must never import it).
"""

from __future__ import annotations

import base64
import logging
import re
from decimal import Decimal
from io import BytesIO
from uuid import UUID, uuid4

import anthropic
from docflow_core import (
    document_status,
    example_prompting,
    field_schema,
    file_types,
    founder_alerts,
    model_runs,
    previews,
    review_digest,
)
from docflow_core.buyers import identify_and_link_buyer
from docflow_core.config import get_settings
from docflow_core.db import tenant_session
from docflow_core.duplicates import detect_document_relationships
from docflow_core.example_prompting import ExamplePlan
from docflow_core.extraction import (
    ExtractionResult,
    build_text_content,
    extract_document,
    wrap_document_content,
)
from docflow_core.field_schema import FieldSchema
from docflow_core.matching import match_document_lines
from docflow_core.numbers import plain_or_none
from docflow_core.storage import read_file, save_file
from docflow_core.validation import validate_document
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.celery_app import celery_app
from app.conversion import ConversionError, PreparedArtifact, format_cell_value, prepare_artifacts

logger = logging.getLogger(__name__)

# Below this many extracted characters, a PDF is treated as scanned/
# image-only and sent to the model visually instead (mirrors docs/parse_pos.py).
PDF_MIN_TEXT_CHARS = 100

_IMAGE_MEDIA_TYPES = {
    file_types.FileTypeName.PNG: "image/png",
    file_types.FileTypeName.JPEG: "image/jpeg",
    file_types.FileTypeName.GIF: "image/gif",
    file_types.FileTypeName.WEBP: "image/webp",
}

_PLAIN_TEXT_LIKE = {
    file_types.FileTypeName.TXT,
    file_types.FileTypeName.CSV,
    file_types.FileTypeName.MD,
    file_types.FileTypeName.HTML,
}


def _extract_pdf_text(content: bytes) -> str:
    """
    Also the enforcement point for two Section 7.11 limits a PDF can only be
    checked for once it is open: the page-count cap
    (`file_types.MAX_DOCUMENT_PAGES`) and password protection. Both raise a
    catalog-coded ConversionError, so the task records the same clean
    `failed` it records for a Tier 2 conversion failure -- never a crash,
    and the file is never handed to the model.
    """
    import pdfplumber
    from pdfminer.pdfdocument import PDFPasswordIncorrect

    try:
        with pdfplumber.open(BytesIO(content)) as pdf:
            page_count = len(pdf.pages)
            if page_count > file_types.MAX_DOCUMENT_PAGES:
                raise ConversionError(
                    "DOC-016",
                    f"PDF has {page_count} pages (limit {file_types.MAX_DOCUMENT_PAGES}).",
                )
            chunks = [page.extract_text() or "" for page in pdf.pages]
    except PDFPasswordIncorrect as exc:
        # Never brute-forced, never passed to the model (Section 7.11).
        raise ConversionError("DOC-001", "PDF is password-protected.") from exc
    return "\n".join(chunks).strip()


def _extract_docx_text(content: bytes) -> str:
    """
    Paragraphs *and* tables: a purchase order's line items almost always
    live in a table, and python-docx's `paragraphs` skips table content
    entirely. Nothing is executed -- python-docx reads the XML part only, it
    has no macro or embedded-object path (Section 7.11: "No active content,
    ever").
    """
    import docx

    document = docx.Document(BytesIO(content))
    lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.replace("\n", " ").strip() for cell in row.cells]
            if any(cells):
                lines.append(", ".join(cells))
    return "\n".join(lines)


def _extract_xlsx_text(content: bytes) -> str:
    import openpyxl

    workbook = openpyxl.load_workbook(BytesIO(content), data_only=True, keep_vba=False)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                lines.append(", ".join(format_cell_value(cell) for cell in row))
    return "\n".join(lines)


_RTF_HEX_RE = re.compile(r"\\'([0-9a-fA-F]{2})")
_RTF_CONTROL_RE = re.compile(r"\\([a-zA-Z]+)(-?\d+)?[ ]?")
# Destination groups whose contents are metadata, not document text.
_RTF_SKIPPED_DESTINATIONS = {
    "fonttbl",
    "colortbl",
    "stylesheet",
    "info",
    "pict",
    "object",
    "themedata",
    "colorschememapping",
    "latentstyles",
    "datastore",
    "generator",
}


def _extract_rtf_text(content: bytes) -> str:
    """
    A deliberately small, dependency-free RTF reader: it walks the control
    words it needs for text flow and drops every destination group that
    carries metadata, embedded objects or pictures. It is a *reader*, not an
    interpreter -- there is no path here that executes an embedded object,
    which is exactly why RTF is not handed to a heavier library
    (Section 7.11: "No active content, ever").
    """
    text_value = content.decode("cp1252", errors="replace")
    out: list[str] = []
    skip_depth = 0
    depth = 0
    index = 0
    length = len(text_value)
    while index < length:
        char = text_value[index]
        if char == "{":
            depth += 1
            index += 1
            continue
        if char == "}":
            if skip_depth and depth <= skip_depth:
                skip_depth = 0
            depth -= 1
            index += 1
            continue
        if char == "\\":
            hex_match = _RTF_HEX_RE.match(text_value, index)
            if hex_match:
                if not skip_depth:
                    out.append(bytes([int(hex_match.group(1), 16)]).decode("cp1252", errors="replace"))
                index = hex_match.end()
                continue
            if index + 1 < length and text_value[index + 1] in "\\{}":
                if not skip_depth:
                    out.append(text_value[index + 1])
                index += 2
                continue
            control_match = _RTF_CONTROL_RE.match(text_value, index)
            if control_match:
                word = control_match.group(1)
                if word == "*" or word in _RTF_SKIPPED_DESTINATIONS:
                    skip_depth = skip_depth or depth
                elif not skip_depth:
                    if word in ("par", "line", "row", "sect", "page"):
                        out.append("\n")
                    elif word in ("tab", "cell"):
                        out.append("\t")
                index = control_match.end()
                continue
            index += 1
            continue
        if not skip_depth:
            out.append(char)
        index += 1
    lines = [" ".join(line.split()) for line in "".join(out).splitlines()]
    return "\n".join(line for line in lines if line)


class UnhandledFileTypeError(Exception):
    pass


def _inner_blocks(file_type: file_types.FileType, content: bytes) -> list[dict]:
    """
    The content blocks for one Tier 1 artifact, *without* the <document>
    delimiters -- so several artifacts (a multi-page TIFF's pages, an
    Outlook message's body plus its attachment) can be wrapped in one
    envelope by `build_content_blocks_for_artifacts`.
    """
    name = file_type.name

    if name == file_types.FileTypeName.PDF:
        text_content = _extract_pdf_text(content)
        if len(text_content) >= PDF_MIN_TEXT_CHARS:
            return [{"type": "text", "text": text_content}]
        encoded = base64.standard_b64encode(content).decode("utf-8")
        return [
            {
                "type": "document",
                "source": {"type": "base64", "media_type": "application/pdf", "data": encoded},
            }
        ]

    if name in _IMAGE_MEDIA_TYPES:
        encoded = base64.standard_b64encode(content).decode("utf-8")
        return [
            {
                "type": "image",
                "source": {
                    "type": "base64",
                    "media_type": _IMAGE_MEDIA_TYPES[name],
                    "data": encoded,
                },
            }
        ]

    if name in _PLAIN_TEXT_LIKE:
        return [{"type": "text", "text": content.decode("utf-8", errors="replace")}]

    if name == file_types.FileTypeName.RTF:
        return [{"type": "text", "text": _extract_rtf_text(content)}]

    if name == file_types.FileTypeName.DOCX:
        return [{"type": "text", "text": _extract_docx_text(content)}]

    if name == file_types.FileTypeName.XLSX:
        return [{"type": "text", "text": _extract_xlsx_text(content)}]

    raise UnhandledFileTypeError(f"No Tier 1 content-block builder for {name}")


def _envelope(parts: list[dict]) -> list[dict]:
    if len(parts) == 1 and parts[0]["type"] == "text":
        return build_text_content(parts[0]["text"])
    return wrap_document_content(parts)


def build_content_blocks(file_type: file_types.FileType, content: bytes) -> list[dict]:
    """
    Native-text PDFs and text-like formats are sent as text; images and
    text-less (scanned) PDFs are sent to the model visually, matching
    docs/parse_pos.py's proven approach, extended to Tier 1's full format
    list (CLAUDE.md Section 8.2 / DECISIONS.md D-010). Tier 2 formats never
    reach this function as themselves -- they arrive as the Tier 1 artifacts
    app/conversion.py produced from them.
    """
    return _envelope(_inner_blocks(file_type, content))


def _artifact_parts(artifacts: list[PreparedArtifact]) -> list[dict]:
    parts: list[dict] = []
    for artifact in artifacts:
        parts.extend(_inner_blocks(artifact.file_type, artifact.content))
    return parts


def build_content_blocks_for_artifacts(artifacts: list[PreparedArtifact]) -> list[dict]:
    """
    One <document> envelope for the whole document, however many Tier 1
    artifacts it turned into (CLAUDE.md Section 7.2: document content is
    always delimited, and it is data, never instructions).
    """
    return _envelope(_artifact_parts(artifacts))


# Stands in, in a text preview, for a page the model read visually (a scanned
# PDF or an image attached to an email). Deliberately generic: an artifact's
# label can carry an attachment filename, which is untrusted (Section 7.11).
_VISUAL_PART_PLACEHOLDER = (
    "[An attached page or image was read visually and cannot be shown as text here. "
    "Download the original to see it.]"
)


def build_preview(
    file_type: file_types.FileType, content: bytes, parts: list[dict]
) -> previews.Preview | None:
    """
    A viewable rendering for a format no browser displays (DECISIONS.md D-092),
    or None when the browser can show the original as it is.

    Image-like formats (TIFF, HEIC) are re-encoded from the original. For
    every other format the preview is the text this task extracted and sent
    to the model -- one extraction, two uses, so the reviewer sees exactly
    what DocFlow read, tables included, for every format the worker opens.
    """
    name = file_type.name
    if not previews.needs_preview(name):
        return None
    if previews.is_image_like(name):
        return previews.build_preview(content, name)
    chunks = [part["text"] if part["type"] == "text" else _VISUAL_PART_PLACEHOLDER for part in parts]
    return previews.text_preview("\n\n".join(chunks))


_PREVIEW_SUFFIXES = {"image/png": ".png", "image/jpeg": ".jpg"}


def _store_preview(
    tenant_id: UUID, document_id: UUID, file_type: file_types.FileType, content: bytes, parts: list[dict]
) -> None:
    """
    Best effort, by design: a preview is a convenience, and a document that
    cannot be previewed still reviews fine against its extracted values --
    the viewer says so. A failure here never touches the document's status.
    """
    try:
        preview = build_preview(file_type, content, parts)
        if preview is None:
            return
        suffix = _PREVIEW_SUFFIXES.get(preview.media_type, ".txt")
        preview_path = save_file(tenant_id, f"preview{suffix}", preview.content)
        with tenant_session(tenant_id) as session:
            session.execute(
                text(
                    """
                    UPDATE documents
                    SET preview_storage_path = :path, preview_media_type = :media_type,
                        preview_kind = :kind
                    WHERE id = :id
                    """
                ),
                {
                    "id": str(document_id),
                    "path": preview_path,
                    "media_type": preview.media_type,
                    "kind": preview.kind,
                },
            )
    except Exception as exc:  # noqa: BLE001 -- see the docstring
        logger.error("preview_failed document_id=%s error_type=%s", document_id, type(exc).__name__)


def _store_extracted_text(tenant_id: UUID, document_id: UUID, parts: list[dict]) -> None:
    """
    Keep the text this task sent the model, so the order can later be shown
    to the model as a past example (Section 7.13: "the same extracted text
    the parser produced"). Only when every part is text: a document read
    visually has no text of its own and can never be an example -- an example
    is never a file or an image. Best effort, like the preview: a document
    whose text couldn't be kept reviews and exports exactly the same.
    """
    if not parts or any(part["type"] != "text" for part in parts):
        return
    body = "\n\n".join(part["text"] for part in parts).strip()
    if not body:
        return
    try:
        path = save_file(tenant_id, "extracted.txt", body.encode("utf-8"))
        with tenant_session(tenant_id) as session:
            session.execute(
                text("UPDATE documents SET extracted_text_path = :path WHERE id = :id"),
                {"id": str(document_id), "path": path},
            )
    except Exception as exc:  # noqa: BLE001 -- see the docstring
        logger.error(
            "extracted_text_store_failed document_id=%s error_type=%s", document_id, type(exc).__name__
        )


def _plan_examples(
    client: anthropic.Anthropic,
    tenant_id: UUID,
    document_id: UUID,
    sender_email: str | None,
    parts: list[dict],
) -> ExamplePlan:
    """
    Approved-example prompting (Section 7.13). Any failure here means "no
    examples", which is exactly how a tenant with the feature off is read --
    the feature must never cost a document its extraction.
    """
    try:
        plan = example_prompting.plan(
            client,
            tenant_id,
            document_id,
            sender_email=sender_email,
            parts=parts,
            session_factory=tenant_session,
        )
    except Exception as exc:  # noqa: BLE001 -- see the docstring
        logger.error("example_planning_failed document_id=%s error_type=%s", document_id, type(exc).__name__)
        return ExamplePlan(outcome="planning_failed")
    # IDs and outcomes only (Section 7.10).
    logger.info(
        "example_plan document_id=%s outcome=%s identified_by=%s examples=%d routing=%s",
        document_id,
        plan.outcome,
        plan.identified_by,
        len(plan.examples),
        plan.routing is not None,
    )
    return plan


def _record_runs(
    session: Session, tenant_id: UUID, document_id: UUID, plan: ExamplePlan, result: ExtractionResult
) -> UUID:
    """Every model call this document cost, one extraction_runs row each
    (D-142) -- the routing read first, because it happened first."""
    if plan.routing is not None:
        model_runs.record_routing(session, tenant_id, document_id, plan.routing)
    return model_runs.record_extraction(session, tenant_id, document_id, result)


def _total_cost(plan: ExamplePlan, result: ExtractionResult) -> Decimal | None:
    """The document's whole model bill: extraction plus the routing read, so
    cost per document and margin on the dashboard stay honest."""
    routing_cost = plan.routing.est_cost_usd if plan.routing is not None else None
    costs = [c for c in (result.est_cost_usd, routing_cost) if c is not None]
    return sum(costs, Decimal("0")) if costs else None


def _money(value: Decimal | None) -> str | None:
    return str(value) if value is not None else None


def _mark_failed(tenant_id: UUID, document_id: UUID, *, raw_response: dict | None = None) -> None:
    """A failure before the model was asked (an unsafe or unconvertible file).
    Its code lives in raw_json, as since D-145, and in failure_code."""
    code = (raw_response or {}).get("error_code")
    with tenant_session(tenant_id) as session:
        document_status.transition(
            session,
            document_id,
            from_statuses=["processing"],
            to="failed",
            values={"raw_json": raw_response, "failure_code": code, "processed_at": document_status.NOW},
        )
    # A code whose catalog message says "DocFlow has been alerted" (a failed
    # conversion, DOC-017) raises that alert (D-145).
    _alert_failure(tenant_id, document_id, code)


def _fail_after_extraction(
    tenant_id: UUID, document_id: UUID, code: str, *, raw_response: dict | None = None
) -> None:
    """A failure after the model answered (saving it, or checking it: DOC-021).
    The answer already on the document stays; `raw_response` is written only
    when saving it was what failed. Then the promised alert (D-145)."""
    values: dict = {"failure_code": code, "processed_at": document_status.NOW}
    if raw_response is not None:
        values["raw_json"] = raw_response
    with tenant_session(tenant_id) as session:
        document_status.transition(
            session, document_id, from_statuses=["processing"], to="failed", values=values
        )
    _alert_failure(tenant_id, document_id, code)


def _alert_failure(tenant_id: UUID, document_id: UUID, error_code: str | None) -> None:
    """
    Raise the founder alert a failure's catalog message promises (D-145).

    Its own transaction, after the document is already marked failed: an
    alert that can't be written must never roll back the record of what
    happened to the document. It is logged instead, with IDs only.
    """
    try:
        with tenant_session(tenant_id) as session:
            founder_alerts.raise_for_failure(
                session, tenant_id=tenant_id, error_code=error_code, document_id=document_id
            )
    except Exception:
        logger.exception(
            "founder_alert_not_raised document_id=%s error_code=%s", document_id, error_code
        )


_PROVENANCE_HEADER_FIELDS = (
    "po_number",
    "order_date",
    "requested_delivery_date",
    "buyer_name",
    "buyer_contact_email",
    "ship_to_address",
    "payment_terms",
    "order_total",
    "currency",
    "notes",
)

_PROVENANCE_LINE_FIELDS = ("sku", "description", "quantity", "unit", "unit_price", "line_total")


def _extracted_provenance(values: dict, fields: tuple[str, ...]) -> dict:
    """
    Per-field provenance (Section 9 / Section 7.1: "Provenance for every
    value: extracted, edited-by-human, or mapped"). Every value this task
    writes came straight from the model, so every present field is
    `extracted`; a field the model returned as null has no value and
    therefore no provenance. `learned_rule:<id>` and
    `human_edit:<review_action_id>` are written by the matching and review
    slices that produce them.
    """
    return {name: "extracted" for name in fields if values.get(name) is not None}


def _overall_confidence(header_confidence: dict, schema: FieldSchema) -> Decimal:
    """
    CLAUDE.md Section 7.1: "Overall document confidence is the minimum of
    required-field confidences, not an average." Required means required for
    THIS tenant (D-120): the fields its current field schema marks required.
    """
    return schema.overall_confidence(header_confidence)


@celery_app.task(name="docflow.parse_and_extract")
def parse_and_extract(tenant_id: str, document_id: str) -> None:
    """
    Read one document and put it in front of a reviewer -- once, and only
    when it has been fully checked (review findings H1, H3; D-158).

    * **Claim first.** The document moves to `processing` by compare-and-set
      (`document_status.claim_for_processing`). A redelivered or duplicated
      job, or one for a document that has since been approved, finds nothing
      to claim and does nothing -- the queue acknowledges late, so a worker
      restart redelivers, and an approved order must never be knocked back.
    * **The model's answer is saved while still `processing`**, in one
      transaction with the header, the lines and the run record. If the
      worker dies after that, the next claim resumes from the saved answer
      and never pays for a second extraction.
    * **`needs_review` only with the checks.** Buyer identification, matching
      and duplicate detection each run in their own transaction (a failure
      there must never cost the paid-for extraction, D-058) and a step that
      fails is recorded, not just logged. Validation then runs, and its
      warnings -- including VAL-016 for any step that didn't finish -- are
      committed in the SAME transaction as the move to `needs_review`. If
      validation itself fails, the document is `failed` with DOC-021 and the
      founder is alerted: never in review with zero warnings.
    """
    tid = UUID(tenant_id)
    did = UUID(document_id)

    with tenant_session(tid) as session:
        row = session.execute(
            text(
                "SELECT storage_path, original_filename, status, sender_email, "
                "current_extraction_run_id FROM documents WHERE id = :id"
            ),
            {"id": str(did)},
        ).mappings().first()
        if row is None:
            logger.error("parse_and_extract_missing_document document_id=%s", did)
            return
        # A test-batch file (`staged`) waits for the founder's "Run
        # extraction" (D-112), a held one (`quarantined`) for its release,
        # and anything already processed is done: none of them is claimable,
        # so none of them reaches the model from here.
        if not document_status.claim_for_processing(session, did):
            logger.info(
                "parse_and_extract_not_claimed document_id=%s status=%s", did, row["status"]
            )
            return
        storage_path = row["storage_path"]
        original_filename = row["original_filename"]
        sender_email = row.get("sender_email")
        already_extracted = row.get("current_extraction_run_id") is not None

    if already_extracted:
        # An earlier attempt saved the model's answer and then stopped (a
        # restart, a crash, a timeout). Resume from it: no second model call.
        logger.info("parse_and_extract_resuming document_id=%s", did)
        _finish(tid, did)
        return

    content = read_file(storage_path)

    # Defense in depth (CLAUDE.md Section 7.11): re-validate here too. Never
    # trust that upload-time validation still holds by the time this task
    # runs -- the file on disk is the same hostile input either way.
    validation = file_types.validate_upload(content, original_filename)
    if not validation.ok:
        logger.error(
            "parse_and_extract_revalidation_failed document_id=%s error_code=%s",
            did,
            validation.error_code,
        )
        _mark_failed(
            tid, did, raw_response={"error_code": validation.error_code, "detail": validation.detail}
        )
        return

    # Tier 2 conversion and .msg/.eml unwrapping happen here, in the isolated
    # worker, before any parsing (CLAUDE.md Section 7.11). Every converted
    # artifact is re-validated inside prepare_artifacts; a conversion failure
    # is a clean, catalog-coded `failed`, never a crash.
    try:
        artifacts = prepare_artifacts(validation.file_type, content)
        parts = _artifact_parts(artifacts)
        content_blocks = _envelope(parts)
    except ConversionError as exc:
        logger.error(
            "parse_and_extract_conversion_failed document_id=%s error_code=%s",
            did,
            exc.error_code,
        )
        _mark_failed(tid, did, raw_response={"error_code": exc.error_code, "detail": exc.detail})
        return
    except Exception:
        logger.exception("parse_and_extract_parse_error document_id=%s", did)
        _mark_failed(tid, did, raw_response={"error_code": "DOC-005", "detail": "Parsing failed."})
        return

    # Before the model call, so the reviewer can see the original even if
    # extraction fails. Stays here in the worker: building a preview is
    # parsing (Section 7.11), and the API only serves what this wrote.
    _store_preview(tid, did, validation.file_type, content, parts)
    _store_extracted_text(tid, did, parts)

    settings = get_settings()
    client = anthropic.Anthropic(api_key=settings.anthropic_api_key)
    plan = _plan_examples(client, tid, did, sender_email, parts)
    result = extract_document(client, content_blocks, examples=plan.examples)

    if not result.ok:
        with tenant_session(tid) as session:
            _record_runs(session, tid, did, plan, result)
            document_status.transition(
                session,
                did,
                from_statuses=["processing"],
                to="failed",
                values={
                    "model_id": result.model_id,
                    "prompt_hash": result.prompt_hash,
                    "schema_version": result.schema_version,
                    "raw_json": result.raw_response,
                    "est_cost_usd": _money(_total_cost(plan, result)),
                    "failure_code": result.error_code or "DOC-008",
                    "processed_at": document_status.NOW,
                },
            )
        # DOC-008 / DOC-009 / DOC-020 tell the reader DocFlow has been alerted (D-145).
        _alert_failure(tid, did, result.error_code or "DOC-008")
        return

    try:
        _save_extraction(tid, did, plan, result)
    except Exception as exc:  # noqa: BLE001 -- the document must not be stranded (M3)
        # Nothing of the answer was written (one transaction). Keep the paid
        # answer on the failed document, with a code, and tell the founder.
        logger.error("extraction_save_failed document_id=%s error_type=%s", did, type(exc).__name__)
        _fail_after_extraction(tid, did, "DOC-021", raw_response=result.raw_response)
        return

    _finish(tid, did)


def _save_extraction(tid: UUID, did: UUID, plan: ExamplePlan, result: ExtractionResult) -> None:
    """The model's answer, the header, the lines and the run record, in one
    transaction. The status stays `processing`: nothing is reviewable yet."""
    header = result.header
    with tenant_session(tid) as session:
        # The tenant's field schema decides which fields count towards the
        # document's confidence and which are checked (D-120). Recorded on the
        # document, so a later change never makes its numbers unexplainable.
        schema = field_schema.current(session, tid)
        overall_confidence = _overall_confidence(result.header_confidence, schema)
        run_id = _record_runs(session, tid, did, plan, result)
        session.execute(
            text(
                """
                UPDATE documents
                SET model_id = :model_id, prompt_hash = :prompt_hash,
                    schema_version = :schema_version, input_tokens = :input_tokens,
                    output_tokens = :output_tokens, est_cost_usd = :est_cost_usd,
                    injection_suspected = :injection_suspected,
                    overall_confidence = :overall_confidence, raw_json = :raw_json,
                    field_schema_version = :field_schema_version,
                    current_extraction_run_id = :run_id
                WHERE id = :id AND status = 'processing'
                """
            ),
            {
                "id": str(did),
                "model_id": result.model_id,
                "prompt_hash": result.prompt_hash,
                "schema_version": result.schema_version,
                "field_schema_version": schema.version or None,
                "input_tokens": result.input_tokens,
                "output_tokens": result.output_tokens,
                "est_cost_usd": _money(_total_cost(plan, result)),
                "run_id": str(run_id),
                "injection_suspected": result.injection_suspected,
                "overall_confidence": str(overall_confidence),
                "raw_json": result.raw_response,
            },
        )

        session.execute(
            text(
                """
                INSERT INTO document_headers
                    (document_id, tenant_id, po_number, order_date, requested_delivery_date,
                     buyer_name, buyer_contact_email, ship_to_address, payment_terms,
                     order_total, currency, notes, header_confidence, currency_inferred,
                     field_provenance)
                VALUES
                    (:document_id, :tenant_id, :po_number, :order_date, :requested_delivery_date,
                     :buyer_name, :buyer_contact_email, :ship_to_address, :payment_terms,
                     :order_total, :currency, :notes, :header_confidence, :currency_inferred,
                     :field_provenance)
                """
            ),
            {
                "document_id": str(did),
                "tenant_id": str(tid),
                "po_number": header["po_number"],
                "order_date": header["order_date"],
                "requested_delivery_date": header["requested_delivery_date"],
                "buyer_name": header["buyer_name"],
                "buyer_contact_email": header["buyer_contact_email"],
                "ship_to_address": header["ship_to_address"],
                "payment_terms": header["payment_terms"],
                "order_total": plain_or_none(header["order_total"]),
                "currency": header["currency"],
                "notes": header["notes"],
                "header_confidence": result.header_confidence,
                "currency_inferred": result.currency_inferred,
                "field_provenance": _extracted_provenance(header, _PROVENANCE_HEADER_FIELDS),
            },
        )

        for line in result.lines:
            session.execute(
                text(
                    """
                    INSERT INTO document_lines
                        (id, document_id, tenant_id, line_number, sku, description,
                         quantity, unit, unit_price, line_total, confidence, field_provenance)
                    VALUES
                        (:id, :document_id, :tenant_id, :line_number, :sku, :description,
                         :quantity, :unit, :unit_price, :line_total, :confidence, :field_provenance)
                    """
                ),
                {
                    "id": str(uuid4()),
                    "document_id": str(did),
                    "tenant_id": str(tid),
                    "line_number": line["line_number"],
                    "sku": line["sku"],
                    "description": line["description"],
                    "quantity": plain_or_none(line["quantity"]),
                    "unit": line["unit"],
                    "unit_price": plain_or_none(line["unit_price"]),
                    "line_total": plain_or_none(line["line_total"]),
                    "confidence": str(line["confidence"]) if line["confidence"] is not None else None,
                    "field_provenance": _extracted_provenance(line, _PROVENANCE_LINE_FIELDS),
                },
            )


def _finish(tid: UUID, did: UUID) -> None:
    """Post-processing, then validation and the move to `needs_review` in one
    transaction. Runs after a fresh extraction and when resuming one."""
    with tenant_session(tid) as session:
        header = session.execute(
            text(
                "SELECT buyer_name, buyer_contact_email FROM document_headers WHERE document_id = :id"
            ),
            {"id": str(did)},
        ).mappings().first()
        field_schema_version = session.execute(
            text("SELECT field_schema_version FROM documents WHERE id = :id"), {"id": str(did)}
        ).scalar_one_or_none()
        # The same schema version the confidence was computed with, so one
        # document is never half-checked under two versions (D-120).
        schema = field_schema.at_version(session, tid, field_schema_version)

    issues: list[str] = []

    # Buyer identification runs first, because Section 7.6's learned mappings
    # are scoped to a buyer. Own transaction (D-058): a failure here must never
    # roll back the paid-for extraction -- but it is recorded and shown now
    # (VAL-016), not only logged.
    buyer_id: UUID | None = None
    try:
        with tenant_session(tid) as session:
            buyer_id = identify_and_link_buyer(
                session,
                tid,
                did,
                buyer_name=header.get("buyer_name") if header else None,
                buyer_contact_email=header.get("buyer_contact_email") if header else None,
            ).buyer_id
    except Exception as exc:  # noqa: BLE001 -- recorded as a pipeline issue
        # No exception message: a database error's text can carry bound
        # parameters, which here are customer data (Section 7.10).
        logger.error("buyer_identification_failed document_id=%s error_type=%s", did, type(exc).__name__)
        issues.append("buyer_identification")

    # Matching. `buyer_id` None is not a failure: matching then uses the
    # tenant-wide rules only, never another buyer's.
    try:
        with tenant_session(tid) as session:
            summary = match_document_lines(session, tid, did, buyer_id=buyer_id)
        logger.info(
            "matching_complete document_id=%s considered=%d matched=%d",
            did,
            summary.lines_considered,
            summary.lines_matched,
        )
    except Exception as exc:  # noqa: BLE001 -- recorded as a pipeline issue
        logger.error("matching_failed document_id=%s error_type=%s", did, type(exc).__name__)
        issues.append("matching")

    # Duplicate / change-order detection (Section 7.8). Here as well as at
    # ingest because the PO number the change-order case needs only exists
    # after extraction; re-running the content-hash half is idempotent.
    try:
        with tenant_session(tid) as session:
            relationships = detect_document_relationships(session, tid, did)
        logger.info(
            "duplicate_detection_complete document_id=%s duplicate=%s change_order=%s",
            did,
            relationships.is_possible_duplicate,
            relationships.is_possible_change_order,
        )
    except Exception as exc:  # noqa: BLE001 -- recorded as a pipeline issue
        logger.error("duplicate_detection_failed document_id=%s error_type=%s", did, type(exc).__name__)
        issues.append("duplicate_detection")

    # Validation and the move into review: one transaction. Validation reports
    # on everything above -- values, matches, duplicate flags and any step
    # that didn't finish -- and changes no value (Section 7.7).
    try:
        with tenant_session(tid) as session:
            session.execute(
                text("UPDATE documents SET pipeline_issues = :issues WHERE id = :id"),
                {"id": str(did), "issues": issues},
            )
            checked = validate_document(session, tid, did, schema.rules())
            moved = document_status.transition(
                session,
                did,
                from_statuses=["processing"],
                to="needs_review",
                values={"processed_at": document_status.NOW, "failure_code": document_status.NULL},
            )
            if not moved:
                # Someone else moved it while we worked (the stuck sweep gave
                # up on it, say). Their decision stands; undo ours.
                raise _AlreadyMovedOn()
            if issues:
                founder_alerts.raise_alert(
                    session,
                    alert_type="pipeline_step_failed",
                    severity="high",
                    tenant_id=tid,
                    payload={"document_id": str(did), "steps": issues},
                    dedupe_key=f"pipeline_step_failed:{did}",
                )
        # Counts only -- never a value (Section 7.10).
        logger.info(
            "validation_complete document_id=%s warnings=%d created=%d resolved=%d issues=%d",
            did,
            checked.evaluated,
            checked.created,
            checked.resolved,
            len(issues),
        )
    except _AlreadyMovedOn:
        logger.info("parse_and_extract_moved_on document_id=%s", did)
        return
    except Exception as exc:  # noqa: BLE001 -- never silent: failed + DOC-021 + alert
        logger.error("validation_failed document_id=%s error_type=%s", did, type(exc).__name__)
        _fail_after_extraction(tid, did, "DOC-021")
        return

    # The "needs review" digest (slice 5.8c, D-131). Last, in its own
    # transaction: a document in review must never lose its checks over a
    # notification.
    try:
        with tenant_session(tid) as session:
            review_digest.note_needs_review(session, tid, did)
    except Exception as exc:  # noqa: BLE001 -- a notification, not the document
        logger.error("review_digest_failed document_id=%s error_type=%s", did, type(exc).__name__)


class _AlreadyMovedOn(Exception):
    """The document left `processing` under us; roll back and leave it be."""
