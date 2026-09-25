"""
Unit tests for the Tier 1 content-block builder and confidence rollup used
by the isolated parsing task (CLAUDE.md Section 7.11 / Section 7.1). The
full task (DB writes, real extraction call) is covered by
test_parse_and_extract_integration.py, gated on the documents schema
actually being applied (see conftest.py).
"""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO

import docx
import openpyxl
import pytest
from docflow_core.field_schema import DEFAULT_SCHEMA, _resolve
from docflow_core.file_types import FileType, FileTypeName

from app.tasks.parse_and_extract import (
    UnhandledFileTypeError,
    _overall_confidence,
    build_content_blocks,
)


def _ft(name: FileTypeName) -> FileType:
    return FileType(name=name, tier="tier1", media_type="application/octet-stream")


def test_plain_text_becomes_text_content():
    content = "PO Number: 12345".encode("utf-8")
    blocks = build_content_blocks(_ft(FileTypeName.TXT), content)
    assert blocks[0]["type"] == "text"
    assert "PO Number: 12345" in blocks[0]["text"]
    assert "<document>" in blocks[0]["text"]


def test_image_becomes_image_content_block():
    content = b"\x89PNG\r\n\x1a\n" + b"\x00" * 16
    blocks = build_content_blocks(_ft(FileTypeName.PNG), content)
    image_blocks = [b for b in blocks if b["type"] == "image"]
    assert len(image_blocks) == 1
    assert image_blocks[0]["source"]["media_type"] == "image/png"


def test_docx_extracts_paragraph_text():
    document = docx.Document()
    document.add_paragraph("PO Number: BCH-9999")
    document.add_paragraph("Buyer: Acme Test Distributor")
    buf = BytesIO()
    document.save(buf)

    blocks = build_content_blocks(_ft(FileTypeName.DOCX), buf.getvalue())
    assert blocks[0]["type"] == "text"
    assert "BCH-9999" in blocks[0]["text"]
    assert "Acme Test Distributor" in blocks[0]["text"]


def test_xlsx_extracts_cell_text():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["SKU", "Description", "Qty"])
    sheet.append(["CF-1001", "Colombian Whole Bean 5lb", 12])
    buf = BytesIO()
    workbook.save(buf)

    blocks = build_content_blocks(_ft(FileTypeName.XLSX), buf.getvalue())
    assert blocks[0]["type"] == "text"
    assert "CF-1001" in blocks[0]["text"]
    assert "Colombian Whole Bean 5lb" in blocks[0]["text"]


def test_pdf_with_extractable_text_uses_text_path(monkeypatch):
    import app.tasks.parse_and_extract as mod

    monkeypatch.setattr(mod, "_extract_pdf_text", lambda content: "PO Number: BCH-1234\n" * 20)
    blocks = build_content_blocks(_ft(FileTypeName.PDF), b"%PDF-1.4 fake")
    assert blocks[0]["type"] == "text"
    assert "BCH-1234" in blocks[0]["text"]


def test_pdf_without_extractable_text_uses_visual_path(monkeypatch):
    import app.tasks.parse_and_extract as mod

    monkeypatch.setattr(mod, "_extract_pdf_text", lambda content: "")
    blocks = build_content_blocks(_ft(FileTypeName.PDF), b"%PDF-1.4 fake pdf bytes")
    doc_blocks = [b for b in blocks if b["type"] == "document"]
    assert len(doc_blocks) == 1
    assert doc_blocks[0]["source"]["media_type"] == "application/pdf"


def test_unhandled_file_type_raises():
    with pytest.raises(UnhandledFileTypeError):
        build_content_blocks(_ft(FileTypeName.DOC), b"whatever")


class _FakeResult:
    def __init__(self, row):
        self._row = row

    def mappings(self):
        return self

    def first(self):
        return self._row


class _FakeSession:
    """Just enough SQLAlchemy surface to drive the task without a database."""

    def __init__(self, row):
        self._row = row
        self.statements: list[tuple[str, dict]] = []

    def execute(self, statement, params=None):
        sql = str(statement)
        self.statements.append((sql, params or {}))
        if "tenant_field_schemas" in sql:
            # No saved field schema: the tenant is on DocFlow's defaults,
            # which is what every tenant had before D-120.
            return _FakeResult(None)
        if sql.strip().upper().startswith("SELECT"):
            return _FakeResult(self._row)
        return _FakeResult(None)


def test_conversion_failure_marks_the_document_failed_and_keeps_the_worker_alive(monkeypatch):
    """
    CLAUDE.md Section 7.11: "a conversion failure is a clean `failed` with a
    catalog-coded message, never a crash", and "a worker that dies takes one
    document to `failed`, never the app". A truncated TIFF is the malformed
    Tier 2 file; the task must return normally with a DOC-coded failure
    recorded.
    """
    import contextlib
    from uuid import uuid4

    import app.tasks.parse_and_extract as mod

    session = _FakeSession({"storage_path": "tenants/x/uploads/y.tif", "original_filename": "fax.tif"})

    @contextlib.contextmanager
    def fake_tenant_session(tenant_id):
        yield session

    monkeypatch.setattr(mod, "tenant_session", fake_tenant_session)
    monkeypatch.setattr(mod, "read_file", lambda path: b"II\x2a\x00" + b"\xff" * 512)

    mod.parse_and_extract(str(uuid4()), str(uuid4()))

    failures = [
        params
        for sql, params in session.statements
        if "status='failed'" in sql.replace(" ", "") or "status = 'failed'" in sql
    ]
    assert failures, "the document was not marked failed"
    assert failures[-1]["raw_json"]["error_code"].startswith("DOC-")


def test_a_failure_that_promises_an_alert_raises_it_without_risking_the_failed_status(monkeypatch):
    """
    D-145: DOC-017 tells the customer "DocFlow has already been alerted", so
    the worker must raise that alert -- after the failed status is written,
    and never at its expense: an alert that can't be written is logged.
    """
    import contextlib
    from uuid import uuid4

    import app.tasks.parse_and_extract as mod

    session = _FakeSession({"storage_path": "tenants/x/uploads/y.tif", "original_filename": "fax.tif"})

    @contextlib.contextmanager
    def fake_tenant_session(tenant_id):
        yield session

    calls: list[dict] = []

    def broken_alert(session, **kwargs):
        calls.append(kwargs)
        raise RuntimeError("the alert table is unreachable")

    monkeypatch.setattr(mod, "tenant_session", fake_tenant_session)
    monkeypatch.setattr(mod, "read_file", lambda path: b"II\x2a\x00" + b"\xff" * 512)
    monkeypatch.setattr(mod.founder_alerts, "raise_for_failure", broken_alert)

    tenant_id, document_id = uuid4(), uuid4()
    mod.parse_and_extract(str(tenant_id), str(document_id))

    assert calls, "no alert was attempted"
    assert calls[0]["error_code"].startswith("DOC-")
    assert calls[0]["document_id"] == document_id
    assert [sql for sql, _ in session.statements if "status='failed'" in sql.replace(" ", "")]


def test_provenance_records_extracted_for_every_present_field():
    from app.tasks.parse_and_extract import _PROVENANCE_HEADER_FIELDS, _extracted_provenance

    header = {
        "po_number": "BCH-9999",
        "buyer_name": "Bella's Coffee House",
        "order_date": None,
        "currency": "USD",
    }
    provenance = _extracted_provenance(header, _PROVENANCE_HEADER_FIELDS)

    assert provenance == {"po_number": "extracted", "buyer_name": "extracted", "currency": "extracted"}
    # A field the model returned as null has no value, so it has no provenance.
    assert "order_date" not in provenance


def test_provenance_of_a_document_with_nothing_extracted_is_empty_not_wrong():
    from app.tasks.parse_and_extract import _PROVENANCE_LINE_FIELDS, _extracted_provenance

    assert _extracted_provenance({"sku": None, "description": None}, _PROVENANCE_LINE_FIELDS) == {}


def test_overall_confidence_is_the_minimum_of_required_fields_not_an_average():
    """CLAUDE.md Section 7.1, as the tenant's field schema defines required
    (D-120): an optional field read badly -- or absent -- does not decide how
    much of the order a reviewer should distrust."""
    confidence = {"po_number": 0.98, "buyer_name": 0.9, "order_total": 0.95, "currency": 0.93}
    assert _overall_confidence(confidence, DEFAULT_SCHEMA) == Decimal("0.9")

    with_a_bad_optional = {**confidence, "order_date": 0.4, "payment_terms": 0.0}
    assert _overall_confidence(with_a_bad_optional, DEFAULT_SCHEMA) == Decimal("0.9")

    # Required for this tenant: now it counts.
    stricter = _resolve(2, {"header": {"order_date": "required"}})
    assert _overall_confidence(with_a_bad_optional, stricter) == Decimal("0.4")


def test_overall_confidence_handles_empty():
    assert _overall_confidence({}, DEFAULT_SCHEMA) == Decimal("0")


# ── matching is wired in, and cannot cost the document its extraction ───────


def _no_examples(monkeypatch, mod):
    """The example planner reads the database; these tests have none. A
    tenant with the feature off gets exactly this plan."""
    from docflow_core.example_prompting import ExamplePlan

    monkeypatch.setattr(mod.example_prompting, "plan", lambda *args, **kwargs: ExamplePlan())


def _drive_successful_task(monkeypatch, *, buyer_id, matcher):
    """
    Runs `parse_and_extract` end to end over a plain-text PO with the model
    call, storage and both post-extraction steps replaced. Returns the fake
    session so a test can inspect what was written.
    """
    import contextlib
    from uuid import uuid4

    from docflow_core.buyers import BuyerIdentification
    from docflow_core.extraction import ExtractionResult

    import app.tasks.parse_and_extract as mod

    session = _FakeSession({"storage_path": "tenants/x/uploads/po.txt", "original_filename": "po.txt"})

    @contextlib.contextmanager
    def fake_tenant_session(tenant_id):
        yield session

    header = {key: None for key in mod._PROVENANCE_HEADER_FIELDS}
    header["po_number"] = "BCH-2291"
    header["buyer_name"] = "Bella's Coffee House"
    result = ExtractionResult(
        ok=True,
        model_id="test-model",
        prompt_hash="testhash",
        schema_version="test",
        raw_response={},
        header=header,
        header_confidence={"po_number": 0.97},
        lines=[
            {
                "line_number": 1,
                "sku": "CF-1001",
                "description": "Colombian Whole Bean 5lb",
                "quantity": Decimal("12"),
                "unit": "CS",
                "unit_price": Decimal("47.50"),
                "line_total": Decimal("570.00"),
                "confidence": Decimal("0.95"),
            }
        ],
        injection_suspected=False,
        currency_inferred=False,
    )

    monkeypatch.setattr(mod, "tenant_session", fake_tenant_session)
    monkeypatch.setattr(mod, "read_file", lambda path: b"PURCHASE ORDER\nPO Number: BCH-2291\n")
    monkeypatch.setattr(mod.anthropic, "Anthropic", lambda api_key=None: object())
    monkeypatch.setattr(mod, "extract_document", lambda client, content, **kwargs: result)
    _no_examples(monkeypatch, mod)
    monkeypatch.setattr(mod, "save_file", lambda tenant_id, name, data: f"tenants/{tenant_id}/uploads/x.txt")
    monkeypatch.setattr(
        mod,
        "identify_and_link_buyer",
        lambda *args, **kwargs: BuyerIdentification(buyer_id=buyer_id, created=True),
    )
    monkeypatch.setattr(mod, "match_document_lines", matcher)

    mod.parse_and_extract(str(uuid4()), str(uuid4()))
    return session


def test_matching_runs_after_buyer_identification_with_the_resolved_buyer(monkeypatch):
    """
    Matching is buyer-scoped (CLAUDE.md Section 7.6), so it has to see the
    buyer the previous step resolved -- not re-derive one, and not run first.
    """
    from uuid import uuid4

    from docflow_core.matching import MatchingSummary

    buyer_id = uuid4()
    calls: list[dict] = []

    def matcher(session, tenant_id, document_id, *, buyer_id=None):
        calls.append({"tenant_id": tenant_id, "document_id": document_id, "buyer_id": buyer_id})
        return MatchingSummary(lines_considered=1, lines_matched=1)

    _drive_successful_task(monkeypatch, buyer_id=buyer_id, matcher=matcher)

    assert len(calls) == 1
    assert calls[0]["buyer_id"] == buyer_id


def test_a_matching_failure_leaves_the_extraction_intact(monkeypatch):
    """
    DECISIONS.md D-058, extended to matching: extraction is one model call the
    tenant has already paid for. A failure in the deterministic step that runs
    afterwards must never roll it back or mark the document failed -- the
    document stays `needs_review` with unmatched lines, which is the same
    state a document with no catalog hits legitimately has.
    """
    from uuid import uuid4

    def matcher(session, tenant_id, document_id, *, buyer_id=None):
        raise RuntimeError("catalog unavailable")

    session = _drive_successful_task(monkeypatch, buyer_id=uuid4(), matcher=matcher)

    statuses = [
        params
        for sql, params in session.statements
        if "status = 'needs_review'" in sql or "status='failed'" in sql.replace(" ", "")
    ]
    assert statuses, "the document was never given a final status"
    assert all("raw_json" not in params or params.get("model_id") for params in statuses)
    assert not [sql for sql, _ in session.statements if "status='failed'" in sql.replace(" ", "")]


# ── previews for real intake (DECISIONS.md D-092) ────────────────────────────


def _preview_for(filename: str, content: bytes):
    from docflow_core import file_types

    from app.conversion import prepare_artifacts
    from app.tasks.parse_and_extract import _artifact_parts, build_preview

    validation = file_types.validate_upload(content, filename)
    assert validation.ok, validation.error_code
    parts = _artifact_parts(prepare_artifacts(validation.file_type, content))
    return build_preview(validation.file_type, content, parts)


def test_a_word_preview_includes_the_line_items_table():
    """
    A Word PO's line items live in a table. A preview without them would put
    the extracted lines beside a document that seems not to contain them.
    """
    from tests import fixture_builders as fb

    preview = _preview_for("po.docx", fb.build_docx())

    assert preview is not None and preview.kind == "extracted_text"
    assert b"CF-1001" in preview.content
    assert b"CUP-12" in preview.content


@pytest.mark.parametrize(
    "filename, builder",
    [
        ("po.xlsx", "build_xlsx"),
        ("po.rtf", "build_rtf"),
        ("po.eml", "build_eml"),
        ("po.msg", "build_msg"),
        ("po.xls", "build_xls"),
        ("po.odt", "build_odt"),
        ("po.ods", "build_ods"),
    ],
)
def test_every_text_like_format_the_worker_reads_gets_a_text_preview(filename, builder):
    from tests import fixture_builders as fb

    preview = _preview_for(filename, getattr(fb, builder)())

    assert preview is not None, f"{filename} would show the 'can't display' fallback"
    assert preview.kind == "extracted_text"
    assert preview.media_type.startswith("text/plain")


@pytest.mark.parametrize("filename, builder", [("fax.tif", "build_tiff"), ("photo.heic", "build_heic")])
def test_image_like_formats_get_a_converted_image(filename, builder):
    from tests import fixture_builders as fb

    preview = _preview_for(filename, getattr(fb, builder)())

    assert preview is not None and preview.kind == "converted_image"
    assert preview.media_type == "image/png"


def test_formats_a_browser_shows_get_no_preview():
    from tests import fixture_builders as fb

    assert _preview_for("po.pdf", fb.build_pdf(["PURCHASE ORDER " * 20])) is None
    assert _preview_for("po.txt", b"PURCHASE ORDER\nPO Number: TEST-1\n") is None


def test_a_page_read_visually_is_named_in_the_preview_not_dropped():
    from app.tasks.parse_and_extract import _VISUAL_PART_PLACEHOLDER, build_preview

    docx_type = FileType(name=FileTypeName.DOCX, tier="tier1", media_type="application/octet-stream")
    parts = [{"type": "text", "text": "Order body"}, {"type": "image", "source": {}}]

    preview = build_preview(docx_type, b"", parts)

    assert preview is not None
    assert _VISUAL_PART_PLACEHOLDER.encode() in preview.content


def _drive_task_over(monkeypatch, filename: str, content: bytes, *, save_file):
    import contextlib
    from uuid import uuid4

    from docflow_core.extraction import ExtractionResult

    import app.tasks.parse_and_extract as mod

    session = _FakeSession(
        {"storage_path": f"tenants/x/uploads/y{filename[-5:]}", "original_filename": filename}
    )

    @contextlib.contextmanager
    def fake_tenant_session(tenant_id):
        yield session

    failed = ExtractionResult(
        ok=False, model_id="m", prompt_hash="h", schema_version="s", raw_response={}
    )
    monkeypatch.setattr(mod, "tenant_session", fake_tenant_session)
    monkeypatch.setattr(mod, "read_file", lambda path: content)
    monkeypatch.setattr(mod, "save_file", save_file)
    monkeypatch.setattr(mod.anthropic, "Anthropic", lambda api_key=None: object())
    monkeypatch.setattr(mod, "extract_document", lambda client, blocks, **kwargs: failed)
    _no_examples(monkeypatch, mod)

    tenant_id = uuid4()
    mod.parse_and_extract(str(tenant_id), str(uuid4()))
    return session, tenant_id


def test_the_task_stores_a_preview_under_the_tenant_even_when_extraction_fails(monkeypatch):
    """
    The gap this closes: previews were built only by the demo seed script, so
    a real Word upload showed "can't display". The preview is written before
    the model call, so a reviewer can see a document whose extraction failed.
    """
    from tests import fixture_builders as fb

    saved: list[tuple] = []

    def fake_save(tenant_id, name, data):
        saved.append((tenant_id, name, data))
        return f"tenants/{tenant_id}/uploads/preview.txt"

    session, tenant_id = _drive_task_over(monkeypatch, "po.docx", fb.build_docx(), save_file=fake_save)

    previews_saved = [entry for entry in saved if entry[1].startswith("preview")]
    assert len(previews_saved) == 1 and previews_saved[0][0] == tenant_id
    updates = [params for sql, params in session.statements if "preview_storage_path" in sql]
    assert updates == [
        {
            "id": updates[0]["id"],
            "path": f"tenants/{tenant_id}/uploads/preview.txt",
            "media_type": "text/plain; charset=utf-8",
            "kind": "extracted_text",
        }
    ]


def test_a_preview_failure_never_touches_the_document(monkeypatch):
    from tests import fixture_builders as fb

    def broken_save(tenant_id, name, data):
        raise OSError("disk full")

    session, _ = _drive_task_over(monkeypatch, "po.docx", fb.build_docx(), save_file=broken_save)

    assert not [sql for sql, _ in session.statements if "preview_storage_path" in sql]
    # The run still reached the model call and recorded its outcome.
    assert [params for sql, params in session.statements if params.get("model_id") == "m"]


# ── Section 7.13: the text an example needs, and the runs it costs (D-141/142) ──


def test_the_text_sent_to_the_model_is_kept_for_a_text_document(monkeypatch):
    from tests import fixture_builders as fb

    saved: list[tuple] = []

    def fake_save(tenant_id, name, data):
        saved.append((tenant_id, name, data))
        return f"tenants/{tenant_id}/uploads/{name}"

    session, tenant_id = _drive_task_over(monkeypatch, "po.docx", fb.build_docx(), save_file=fake_save)

    kept = [entry for entry in saved if entry[1] == "extracted.txt"]
    assert len(kept) == 1 and kept[0][0] == tenant_id
    assert b"CF-1001" in kept[0][2]
    updates = [params for sql, params in session.statements if "extracted_text_path" in sql]
    assert updates and updates[0]["path"] == f"tenants/{tenant_id}/uploads/extracted.txt"


def test_a_document_read_visually_keeps_no_text_and_can_never_be_an_example(monkeypatch):
    """An example is never a file or an image (Section 7.13, Section 10)."""
    from tests import fixture_builders as fb

    saved: list[tuple] = []

    def fake_save(tenant_id, name, data):
        saved.append((tenant_id, name, data))
        return f"tenants/{tenant_id}/uploads/{name}"

    session, _ = _drive_task_over(monkeypatch, "fax.tif", fb.build_tiff(), save_file=fake_save)

    assert not [entry for entry in saved if entry[1] == "extracted.txt"]
    assert not [sql for sql, _ in session.statements if "extracted_text_path" in sql]


def test_a_failed_extraction_is_still_recorded_as_a_model_run(monkeypatch):
    """The daily cost breaker reads extraction_runs (Section 7.9); a call that
    failed still happened and must be there."""
    from tests import fixture_builders as fb

    session, _ = _drive_task_over(
        monkeypatch, "po.docx", fb.build_docx(), save_file=lambda t, n, d: f"tenants/{t}/uploads/{n}"
    )

    runs = [params for sql, params in session.statements if "INSERT INTO extraction_runs" in sql]
    assert len(runs) == 1
    assert runs[0]["run_kind"] == "extraction" and runs[0]["succeeded"] is False


def test_a_successful_extraction_records_its_run_and_points_the_document_at_it(monkeypatch):
    session = _drive_successful_task(monkeypatch, buyer_id=None, matcher=lambda *a, **k: _NoMatches())

    runs = [params for sql, params in session.statements if "INSERT INTO extraction_runs" in sql]
    assert len(runs) == 1 and runs[0]["succeeded"] is True and runs[0]["examples_used"] == "[]"
    updates = [params for sql, params in session.statements if "current_extraction_run_id" in sql]
    assert updates and updates[0]["run_id"] == runs[0]["id"]


def test_the_routing_read_is_its_own_run_and_its_cost_is_part_of_the_document(monkeypatch):
    from docflow_core.example_prompting import ExamplePlan
    from docflow_core.extraction import RoutingResult

    import app.tasks.parse_and_extract as mod

    routing = RoutingResult(
        ok=True, model_id="routing-model", prompt_hash="r", schema_version="routing",
        raw_response={}, input_tokens=900, output_tokens=40, est_cost_usd=Decimal("0.0011"),
    )
    captured = {}

    def fake_extract(client, content, **kwargs):
        captured.update(kwargs)
        from docflow_core.extraction import ExtractionResult

        return ExtractionResult(
            ok=False, model_id="m", prompt_hash="h", schema_version="s", raw_response={},
            est_cost_usd=Decimal("0.0100"),
        )

    from tests import fixture_builders as fb

    session, _ = _drive_task_over(
        monkeypatch, "po.docx", fb.build_docx(), save_file=lambda t, n, d: f"tenants/{t}/uploads/{n}"
    )
    # Re-drive with a routing plan in place.
    monkeypatch.setattr(mod.example_prompting, "plan", lambda *a, **k: ExamplePlan(routing=routing))
    monkeypatch.setattr(mod, "extract_document", fake_extract)
    session.statements.clear()
    from uuid import uuid4

    mod.parse_and_extract(str(uuid4()), str(uuid4()))

    runs = [params for sql, params in session.statements if "INSERT INTO extraction_runs" in sql]
    assert [r["run_kind"] for r in runs] == ["buyer_routing", "extraction"]
    assert captured["examples"] == []
    costs = [params["est_cost_usd"] for sql, params in session.statements if "status = 'failed'" in sql]
    assert costs == ["0.0111"]


def test_a_planning_failure_means_no_examples_never_a_lost_document(monkeypatch):
    from uuid import uuid4

    import app.tasks.parse_and_extract as mod

    def broken_plan(*args, **kwargs):
        raise RuntimeError("database away")

    monkeypatch.setattr(mod.example_prompting, "plan", broken_plan)
    plan = mod._plan_examples(object(), uuid4(), uuid4(), None, [])
    assert plan.examples == [] and plan.outcome == "planning_failed"


class _NoMatches:
    lines_considered = 0
    lines_matched = 0
