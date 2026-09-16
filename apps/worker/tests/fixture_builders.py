"""
Programmatic fixture builders for the Section 7.11 required tests.

Everything here is generated at test time rather than committed as a binary
blob, so the repository carries no opaque files and every fixture's exact
content is readable. All sample content is unmistakably fake (CLAUDE.md
Section 0 rule 4): the same fictional "Bella's Coffee House" buyer the
golden fixture in docs/sample_po.txt already uses.

The OLE compound-file writer exists because no Python library writes
`.msg`: without it there would be no way to exercise the `.msg` unwrapping
path with a real file. It writes the minimum valid MS-CFB v3 container --
every stream is padded to at least the 4096-byte mini-stream cutoff so no
mini-FAT is needed, which is legal and keeps the writer small.
"""

from __future__ import annotations

import struct
import zipfile
from dataclasses import dataclass, field
from io import BytesIO

SECTOR_SIZE = 512
DIRECTORY_ENTRY_SIZE = 128
MINI_STREAM_CUTOFF = 4096

FREESECT = 0xFFFFFFFF
ENDOFCHAIN = 0xFFFFFFFE
FATSECT = 0xFFFFFFFD
NOSTREAM = 0xFFFFFFFF

PO_TEXT = (
    "PURCHASE ORDER\n"
    "PO Number: BCH-2291\n"
    "Order Date: 2026-03-14\n"
    "Buyer: Bella's Coffee House\n"
    "Contact: orders@bellascoffee.com\n"
    "Payment Terms: Net 30\n"
    "\n"
    "SKU, Description, Qty, Unit, Unit Price, Line Total\n"
    "CF-1001, Colombian Whole Bean 5lb, 12, CS, 47.50, 570.00\n"
    "CF-2210, Ethiopian Yirgacheffe 5lb, 6, CS, 62.00, 372.00\n"
    "SY-0045, Vanilla Syrup 750ml, 24, EA, 8.25, 198.00\n"
    "CUP-12, 12oz Paper Cups (1000ct), 4, BOX, 54.00, 216.00\n"
    "\n"
    "Order Total: 1356.00\n"
)

PO_ROWS = [
    ["PURCHASE ORDER", "", "", "", "", ""],
    ["PO Number", "BCH-2291", "", "", "", ""],
    ["Buyer", "Bella's Coffee House", "", "", "", ""],
    ["SKU", "Description", "Qty", "Unit", "Unit Price", "Line Total"],
    ["CF-1001", "Colombian Whole Bean 5lb", 12, "CS", 47.50, 570.00],
    ["CF-2210", "Ethiopian Yirgacheffe 5lb", 6, "CS", 62.00, 372.00],
    ["SY-0045", "Vanilla Syrup 750ml", 24, "EA", 8.25, 198.00],
    ["CUP-12", "12oz Paper Cups (1000ct)", 4, "BOX", 54.00, 216.00],
    ["Order Total", "", "", "", "", 1356.00],
]


# ── OLE compound file (.msg) ───────────────────────────────────────────────


@dataclass
class _Entry:
    name: str
    entry_type: int  # 5 root, 1 storage, 2 stream
    data: bytes = b""
    # Padding unit for this stream: a UTF-16LE space for text-typed MAPI
    # properties, a plain space for binary ones, so padding a stream past the
    # mini-stream cutoff never corrupts the value it carries.
    pad_unit: bytes = b" "
    children: list["_Entry"] = field(default_factory=list)
    start_sector: int = ENDOFCHAIN
    size: int = 0
    child_id: int = NOSTREAM
    right_id: int = NOSTREAM


def _pad_stream(data: bytes, pad_unit: bytes) -> bytes:
    """
    Every stream is padded past the mini-stream cutoff so the container needs
    no mini-FAT. The padding is trailing whitespace in the stream's own
    encoding, which every consumer of these fixtures strips.
    """
    if len(data) >= MINI_STREAM_CUTOFF:
        return data
    shortfall = MINI_STREAM_CUTOFF - len(data)
    units = (shortfall + len(pad_unit) - 1) // len(pad_unit)
    return data + pad_unit * units


def _flatten(root: _Entry) -> list[_Entry]:
    """
    Assigns directory ids and links each storage's children as a
    right-leaning chain (a legal, if unbalanced, directory tree).
    """
    ordered: list[_Entry] = [root]

    def visit(entry: _Entry) -> None:
        if not entry.children:
            return
        first_index = len(ordered)
        ordered.extend(entry.children)
        entry.child_id = first_index
        for offset, child in enumerate(entry.children):
            child.right_id = (
                first_index + offset + 1 if offset + 1 < len(entry.children) else NOSTREAM
            )
        for child in entry.children:
            visit(child)

    visit(root)
    return ordered


def _directory_entry_bytes(entry: _Entry) -> bytes:
    name_utf16 = entry.name.encode("utf-16-le") + b"\x00\x00"
    if len(name_utf16) > 64:
        raise ValueError(f"Directory entry name too long: {entry.name!r}")
    raw = bytearray(DIRECTORY_ENTRY_SIZE)
    raw[0 : len(name_utf16)] = name_utf16
    struct.pack_into("<H", raw, 64, len(name_utf16))
    raw[66] = entry.entry_type
    raw[67] = 1  # black
    struct.pack_into("<I", raw, 68, NOSTREAM)  # left sibling
    struct.pack_into("<I", raw, 72, entry.right_id)
    struct.pack_into("<I", raw, 76, entry.child_id)
    struct.pack_into("<I", raw, 116, entry.start_sector)
    struct.pack_into("<Q", raw, 120, entry.size)
    return bytes(raw)


def build_ole_container(root: _Entry) -> bytes:
    entries = _flatten(root)

    # 1. Lay out stream data sectors.
    sector_payloads: list[bytes] = []
    fat: list[int] = []
    for entry in entries:
        if entry.entry_type != 2 or not entry.data:
            continue
        data = _pad_stream(entry.data, entry.pad_unit)
        entry.size = len(data)
        first = len(sector_payloads)
        entry.start_sector = first
        chunk_count = (len(data) + SECTOR_SIZE - 1) // SECTOR_SIZE
        for index in range(chunk_count):
            chunk = data[index * SECTOR_SIZE : (index + 1) * SECTOR_SIZE]
            sector_payloads.append(chunk.ljust(SECTOR_SIZE, b"\x00"))
            fat.append(first + index + 1 if index + 1 < chunk_count else ENDOFCHAIN)

    # 2. Directory sectors. Unused slots in the last sector stay zeroed,
    # which is an "unallocated" directory entry -- never reachable from the
    # root's tree, so no reader ever visits one.
    directory_bytes = b"".join(_directory_entry_bytes(entry) for entry in entries)
    directory_bytes += b"\x00" * ((-len(directory_bytes)) % SECTOR_SIZE)
    directory_first_sector = len(sector_payloads)
    directory_sector_count = len(directory_bytes) // SECTOR_SIZE
    for index in range(directory_sector_count):
        sector_payloads.append(
            directory_bytes[index * SECTOR_SIZE : (index + 1) * SECTOR_SIZE]
        )
        fat.append(
            directory_first_sector + index + 1
            if index + 1 < directory_sector_count
            else ENDOFCHAIN
        )

    # 3. FAT sectors, sized to cover themselves.
    entries_per_fat_sector = SECTOR_SIZE // 4
    fat_sector_count = 1
    while True:
        total_sectors = len(sector_payloads) + fat_sector_count
        needed = (total_sectors + entries_per_fat_sector - 1) // entries_per_fat_sector
        if needed <= fat_sector_count:
            break
        fat_sector_count = needed

    fat_first_sector = len(sector_payloads)
    for index in range(fat_sector_count):
        fat.append(FATSECT)
    fat.extend([FREESECT] * (fat_sector_count * entries_per_fat_sector - len(fat)))

    fat_bytes = b"".join(struct.pack("<I", value) for value in fat)
    for index in range(fat_sector_count):
        sector_payloads.append(fat_bytes[index * SECTOR_SIZE : (index + 1) * SECTOR_SIZE])

    # 4. Header.
    header = bytearray(SECTOR_SIZE)
    header[0:8] = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    struct.pack_into("<H", header, 24, 0x003E)  # minor version
    struct.pack_into("<H", header, 26, 0x0003)  # major version 3 (512-byte sectors)
    struct.pack_into("<H", header, 28, 0xFFFE)  # little-endian
    struct.pack_into("<H", header, 30, 9)  # sector shift -> 512
    struct.pack_into("<H", header, 32, 6)  # mini sector shift -> 64
    struct.pack_into("<I", header, 44, fat_sector_count)
    struct.pack_into("<I", header, 48, directory_first_sector)
    struct.pack_into("<I", header, 56, MINI_STREAM_CUTOFF)
    struct.pack_into("<I", header, 60, ENDOFCHAIN)  # first mini-FAT sector
    struct.pack_into("<I", header, 64, 0)  # mini-FAT sector count
    struct.pack_into("<I", header, 68, ENDOFCHAIN)  # first DIFAT sector
    struct.pack_into("<I", header, 72, 0)  # DIFAT sector count
    for slot in range(109):
        value = fat_first_sector + slot if slot < fat_sector_count else FREESECT
        struct.pack_into("<I", header, 76 + slot * 4, value)

    return bytes(header) + b"".join(sector_payloads)


def _msg_text_stream(prop_id: str, value: str) -> _Entry:
    return _Entry(
        name=f"__substg1.0_{prop_id}001F",
        entry_type=2,
        data=value.encode("utf-16-le"),
        pad_unit=" ".encode("utf-16-le"),
    )


def build_msg(
    *,
    subject: str = "PO BCH-2291 from Bella's Coffee House",
    sender: str = "orders@bellascoffee.com",
    body: str = PO_TEXT,
    attachments: list[tuple[str, bytes]] | None = None,
) -> bytes:
    """A minimal but genuine Outlook `.msg` (OLE compound file)."""
    children = [
        _msg_text_stream("0037", subject),
        _msg_text_stream("1000", body),
        _msg_text_stream("0C1F", sender),
    ]
    for index, (filename, data) in enumerate(attachments or []):
        children.append(
            _Entry(
                name=f"__attach_version1.0_#{index:08X}",
                entry_type=1,
                children=[
                    _msg_text_stream("3707", filename),
                    _Entry(name="__substg1.0_37010102", entry_type=2, data=data),
                ],
            )
        )
    root = _Entry(name="Root Entry", entry_type=5, children=children)
    return build_ole_container(root)


def build_ole_with_stream(stream_name: str, data: bytes = b"fake legacy office content") -> bytes:
    """
    An OLE container carrying one named stream -- enough for the allowlist's
    content-level `.doc`/`.xls` classification, without pretending to be a
    parseable legacy document.
    """
    root = _Entry(
        name="Root Entry",
        entry_type=5,
        children=[_Entry(name=stream_name, entry_type=2, data=data)],
    )
    return build_ole_container(root)


# ── PDF ────────────────────────────────────────────────────────────────────


def build_pdf(page_texts: list[str]) -> bytes:
    """
    A minimal, uncompressed, multi-page PDF. Hand-assembled so the page count
    is exactly what the test asks for (a 500-page PDF is one of Section
    7.11's required rejection fixtures) with no dependency on a PDF writer.
    """
    objects: list[bytes] = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    font_id = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    page_ids: list[int] = []
    content_ids: list[int] = []
    for text in page_texts:
        escaped = text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1")
        content_ids.append(add(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"))

    pages_id = len(objects) + len(page_texts) + 1
    for content_id in content_ids:
        page_ids.append(
            add(
                b"<< /Type /Page /Parent %d 0 R /MediaBox [0 0 612 792] "
                b"/Resources << /Font << /F1 %d 0 R >> >> /Contents %d 0 R >>"
                % (pages_id, font_id, content_id)
            )
        )
    kids = b" ".join(b"%d 0 R" % page_id for page_id in page_ids)
    actual_pages_id = add(
        b"<< /Type /Pages /Kids [%s] /Count %d >>" % (kids, len(page_ids))
    )
    assert actual_pages_id == pages_id, "page-tree object id prediction is wrong"
    catalog_id = add(b"<< /Type /Catalog /Pages %d 0 R >>" % pages_id)

    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % index + body + b"\nendobj\n"
    xref_offset = len(out)
    out += b"xref\n0 %d\n" % (len(objects) + 1)
    out += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root %d 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        catalog_id,
        xref_offset,
    )
    return bytes(out)


def build_pdf_with_encrypt_trailer() -> bytes:
    """
    A PDF whose trailer references an encryption dictionary -- the shape a
    password-protected PDF has, which the allowlist detects without opening
    the document (CLAUDE.md Section 7.11: never brute-forced, never passed to
    the model).
    """
    pdf = build_pdf(["Protected purchase order"])
    return pdf.replace(b"trailer\n<< /Size", b"trailer\n<< /Encrypt 99 0 R /Size", 1)


# ── OpenDocument ───────────────────────────────────────────────────────────

_ODT_CONTENT = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content
  xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
  xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
  xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0">
  <office:body><office:text>
    <text:p>PURCHASE ORDER</text:p>
    <text:p>PO Number: BCH-2291</text:p>
    <text:p>Buyer: Bella's Coffee House</text:p>
    <table:table>
      {rows}
    </table:table>
  </office:text></office:body>
</office:document-content>
"""

_ODS_CONTENT = _ODT_CONTENT.replace("office:text", "office:spreadsheet")


def _odf_cell_text(value: object) -> str:
    return f"{value:.2f}" if isinstance(value, float) else str(value)


def _odf_rows() -> str:
    rows = []
    for row in PO_ROWS[3:]:
        cells = "".join(
            f"<table:table-cell><text:p>{_odf_cell_text(value)}</text:p></table:table-cell>"
            for value in row
        )
        rows.append(f"<table:table-row>{cells}</table:table-row>")
    return "".join(rows)


def build_odf(mimetype: str, body_template: str) -> bytes:
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("mimetype", mimetype)
        archive.writestr("content.xml", body_template.format(rows=_odf_rows()))
    return buffer.getvalue()


def build_odt() -> bytes:
    return build_odf("application/vnd.oasis.opendocument.text", _ODT_CONTENT)


def build_ods() -> bytes:
    return build_odf("application/vnd.oasis.opendocument.spreadsheet", _ODS_CONTENT)


# ── Images ─────────────────────────────────────────────────────────────────


def _po_image(width: int = 600, height: int = 400):
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (width, height), (255, 255, 255))
    draw = ImageDraw.Draw(image)
    for index, line in enumerate(PO_TEXT.splitlines()[:12]):
        draw.text((10, 10 + index * 14), line, fill=(0, 0, 0))
    return image


def build_tiff(pages: int = 1, size: tuple[int, int] = (600, 400)) -> bytes:
    images = [_po_image(*size) for _ in range(pages)]
    buffer = BytesIO()
    images[0].save(buffer, format="TIFF", save_all=True, append_images=images[1:])
    return buffer.getvalue()


def build_heic() -> bytes:
    import pillow_heif

    pillow_heif.register_heif_opener()
    buffer = BytesIO()
    _po_image().save(buffer, format="HEIF")
    return buffer.getvalue()


# ── Office ─────────────────────────────────────────────────────────────────


def build_xls() -> bytes:
    import xlwt

    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("PO")
    for row_index, row in enumerate(PO_ROWS):
        for column_index, value in enumerate(row):
            sheet.write(row_index, column_index, value)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_docx() -> bytes:
    import docx

    document = docx.Document()
    document.add_paragraph("PURCHASE ORDER")
    document.add_paragraph("PO Number: BCH-2291")
    document.add_paragraph("Buyer: Bella's Coffee House")
    table = document.add_table(rows=0, cols=6)
    for row in PO_ROWS[3:]:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_xlsx() -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row in PO_ROWS:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_rtf() -> bytes:
    lines = "\\par\n".join(line for line in PO_TEXT.splitlines() if line)
    return (
        "{\\rtf1\\ansi\\deff0{\\fonttbl{\\f0 Helvetica;}}\n" + lines + "\n}"
    ).encode("cp1252")


def build_eml(attachments: list[tuple[str, str, bytes]] | None = None) -> bytes:
    """attachments: (filename, subtype, data) -- subtype is the MIME subtype."""
    from email.message import EmailMessage

    message = EmailMessage()
    message["Subject"] = "PO BCH-2291"
    message["From"] = "orders@bellascoffee.com"
    message["To"] = "orders@acmetestdistributor.example"
    message["Date"] = "Sat, 14 Mar 2026 09:00:00 -0700"
    message.set_content(PO_TEXT)
    for filename, subtype, data in attachments or []:
        message.add_attachment(
            data, maintype="application", subtype=subtype, filename=filename
        )
    return message.as_bytes()
